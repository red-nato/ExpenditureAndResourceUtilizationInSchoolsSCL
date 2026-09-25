"""Genera informe HTML autónomo, Excel y notebook ejecutado sin servidor Jupyter."""
from pathlib import Path
import base64, html, json, os, re, sys
import pandas as pd
import markdown
import nbformat
from IPython.core.interactiveshell import InteractiveShell
from IPython.utils.capture import capture_output
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
P=ROOT/'datos/procesados'
REPORT=ROOT/'informes'

def html_report():
    src=(REPORT/'EDA_gasto_educativo.md').read_text()
    content=markdown.markdown(src,extensions=['tables','fenced_code','toc'])
    def embed(m):
        path=REPORT/m.group(1)
        return f'src="data:image/png;base64,{base64.b64encode(path.read_bytes()).decode()}"'
    content=re.sub(r'src="(graficos/[^"]+)"',embed,content)
    css='''
    :root{color-scheme:light;--ink:#172c39;--muted:#5c6c77;--accent:#166382;--line:#dae1e5}
    *{box-sizing:border-box}body{margin:0;background:#edf1f3;color:var(--ink);font:17px/1.7 system-ui,-apple-system,sans-serif}
    main{max-width:1120px;margin:36px auto;background:#fff;padding:55px 70px;box-shadow:0 8px 36px #1232;border-top:7px solid var(--accent)}
    h1{font-size:42px;line-height:1.16;letter-spacing:-1.2px;max-width:900px;margin:0 0 22px}
    h2{font-size:27px;line-height:1.3;border-top:1px solid var(--line);padding-top:28px;margin-top:42px}
    h3{font-size:20px;margin-top:30px}p{margin:17px 0}a{color:var(--accent);text-underline-offset:3px}
    table{border-collapse:collapse;font-size:14px;line-height:1.5;width:100%;margin:25px 0;display:block;overflow:auto}
    th{background:#193d50;color:white;text-align:left}td,th{padding:12px 14px;border-bottom:1px solid var(--line);vertical-align:top}
    tr:nth-child(even) td{background:#f2f6f7}img{width:100%;height:auto;display:block;margin:24px auto}
    code{font-size:.9em;background:#eef3f5;padding:2px 4px;border-radius:3px}li{margin:10px 0}
    footer{border-top:1px solid var(--line);margin-top:35px;padding-top:20px;font-size:13px;color:var(--muted)}
    @media(max-width:700px){main{margin:0;padding:30px 20px}h1{font-size:32px}body{font-size:16px}}
    @media print{body{background:white;font-size:11pt}main{margin:0;box-shadow:none;padding:0;max-width:none}h1{font-size:28pt}h2{break-after:avoid}img,table{break-inside:avoid}a{color:inherit}}
    '''
    doc=f'<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>EDA · Recursos educativos</title><style>{css}</style></head><body><main>{content}<footer>Datos originales conservados · Cálculos reproducibles · Los valores PME son estimaciones declaradas.</footer></main></body></html>'
    (REPORT/'EDA_gasto_educativo.html').write_text(doc)

def workbook():
    files={
      'Compras CGR':'cgr_compras_lotes','No habidos CGR':'cgr_no_habidos','Con comodato CGR':'cgr_con_comodato',
      'Cruce comodatos':'cgr_interseccion_comodato_no_habidos','Fuentes fondos':'cgr_fuentes_financiamiento',
      'CGR por modelo':'cgr_resumen_modelo','CGR por destino':'cgr_resumen_dependencia','Limites destino':'cgr_limites_por_dependencia',
      'CGR por año':'cgr_resumen_anio','Documentos pago':'cgr_documentos_pago',
      'Panel PME nacional':'pme_panel_establecimientos_2024','Panel pilotos':'pme_panel_pilotos_2024',
      'Sensibilidad PME':'pme_sensibilidad_montos','PME por dimension':'pme_resumen_dimensiones_2024','Correlaciones':'pme_correlaciones','Correlaciones por grupo':'pme_correlaciones_estratificadas',
      'Cobertura':'pme_cobertura_dependencia','Faltantes':'pme_faltantes','Descriptivos':'pme_descriptivos_montos',
      'Niveles implementacion':'pme_nivel_implementacion','Extremos por verificar':'pme_montos_extremos_para_verificar','Controles':'verificaciones'}
    with pd.ExcelWriter(REPORT/'Tablas_EDA.xlsx',engine='openpyxl') as writer:
        notes=pd.DataFrame([
          ['Alcance','Contexto PME 2024; caso Melipeuco compras 2020–2022, inspección 2023, informe 2024.'],
          ['No habidos','Valor histórico observado, no pérdida definitiva ni situación actual verificada.'],
          ['PME','Estimaciones de acciones; no pagos, rendiciones aceptadas ni resultados educativos.'],
          ['Avance','pct_completas es fracción 0–1; avance_min/max son límites 0–100 del promedio simple.'],
          ['Sin dato','Celdas vacías permanecen sin dato. No reemplazar por cero.'],
          ['Filas','id_fila o fila_excel y pagina_pdf permiten revisar el origen.'],
          ['Originales','Consultar datos/fuentes.json para URL, fecha y hash.'],
          ['Diferencia $1','Compras suman $176.164.030; el total impreso es $176.164.029.'],
          ['No aditividad','61 duplicidades TIC ya incluidas en no habidos. Máximos de destino no se suman.'],
          ['Microdatos nacionales','La totalidad de acciones se conserva en CSV para mantener este libro manejable.']
        ],columns=['Concepto','Lectura'])
        notes.to_excel(writer,sheet_name='Leer primero',index=False)
        for sheet,filename in files.items():
            df=pd.read_csv(P/f'{filename}.csv',dtype={'serie':'str','serie_original':'str','factura':'str','decreto':'str'})
            df.to_excel(writer,sheet_name=sheet,index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            ws.row_dimensions[1].height=32
            for cell in ws[1]:
                cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='193D50');cell.alignment=Alignment(wrap_text=True,vertical='center')
            for col in ws.iter_cols():
                label=str(col[0].value)
                width=min(43,max(15,len(label)+3))
                if 'nombre' in label.lower() or 'NOM_RBD'==label or 'dependencia' in label or 'familia'==label:width=33
                ws.column_dimensions[get_column_letter(col[0].column)].width=width
                for cell in col[1:]:
                    if isinstance(cell.value,(float,int)):
                        if label in ['pct_completas','tasa_min','tasa_max','tasa_no_habidos','cobertura','fraccion_completas','fraccion_monto_en_completas']:cell.number_format='0.00%'
                        elif 'clp' in label or label in ['estim_total','estim_sep','estim_por_alumno'] or label.startswith('ESTIM_'):cell.number_format='"$"#,##0'
                        elif label in ['pearson','spearman'] or label.startswith('rho_'):cell.number_format='0.000'
            if ws.title=='Leer primero':ws.column_dimensions['B'].width=110

def notebook():
    md=nbformat.v4.new_markdown_cell;code=nbformat.v4.new_code_cell
    cells=[
      md('# EDA: del gasto educativo a la disponibilidad\n\nDatos reales de Mineduc y Contraloría. **Pregunta:** ¿qué brechas de disponibilidad pueden comprobarse y qué indicadores de planificación son confiables?\n\nCaso principal: Melipeuco. Comparación: PME nacional y La Cisterna, 2024. Este cuaderno no estima una pérdida definitiva ni efectos causales. Las salidas se ejecutan con Python, sin servidor Jupyter; las celdas son reejecutables en Jupyter.'),
      code("from pathlib import Path\nimport sys, json\nimport pandas as pd\nfrom IPython.display import display, Image\nROOT = Path.cwd()\nif not (ROOT / 'scripts').exists(): ROOT = ROOT.parent\nsys.path.insert(0, str(ROOT / 'scripts'))\nP = ROOT / 'datos/procesados'\ndef tabla(nombre): return pd.read_csv(P / (nombre + '.csv'))\nprint('Directorio de trabajo:', ROOT.name)"),
      md('## 1. Descarga, procedencia y limpieza\n\nLos originales se descargan con `scripts/descargar.py` (ver README). Se verifica la versión mediante SHA-256. No se usa información personal de estudiantes en el análisis. La extracción separa tablas que comparten página y preserva los equipos sin serie.'),
      code("fuentes = json.loads((ROOT/'datos/fuentes.json').read_text())\ndisplay(pd.DataFrame(fuentes)[['productor','titulo','periodo','bytes']])"),
      code("from extraer_fuentes import main as extraer\nfrom analizar import analyze_case, analyze_pme\nextraer()\nresultados = {'caso': analyze_case(), 'pme': analyze_pme()}\nprint(json.dumps(resultados, ensure_ascii=False, indent=2))"),
      md('## 2. Calidad y cobertura\n\nSe preservan 4 duplicados PME adicionales, 38 montos SEP ausentes y los extremos. Los faltantes no se convierten en cero. Las correlaciones usan 8.237 escuelas activas con matrícula positiva. `MATRICULA` es una bandera; el denominador correcto es `MAT_TOTAL`.'),
      code("display(tabla('pme_faltantes').query('faltantes > 0').head(15))\ndisplay(tabla('pme_cobertura_dependencia'))\ndisplay(tabla('pme_descriptivos_montos'))"),
      md('## 3. Distribuciones, densidad y sensibilidad\n\nHistograma original y densidad en logaritmos sólo para valores positivos. El 21,25% de acciones tiene monto cero. 19 registros explican 12,80% del total; excluirlos es sensibilidad, no una corrección demostrada. La comparación por dimensión cuenta acciones al 100% y usa estimaciones, no pagos.'),
      code("display(Image(filename=str(ROOT/'informes/graficos/05_pme_histograma_densidad.png')))\ndisplay(tabla('pme_sensibilidad_montos'))\ndisplay(tabla('pme_resumen_dimensiones_2024'))\ndisplay(Image(filename=str(ROOT/'informes/graficos/08_pme_dimensiones.png')))\ndisplay(Image(filename=str(ROOT/'informes/graficos/10_pme_dimensiones.png')))"),
      md('## 4. Correlaciones a la unidad correcta\n\nUna fila por RBD. Spearman ofrece una lectura menos dependiente de los extremos que Pearson. El gasto aquí es estimado. La asociación con implementación declarada es débil; no equivale a eficiencia ni a aprendizaje.'),
      code("c = tabla('pme_correlaciones')\ndisplay(c[(c.variable_1=='MAT_TOTAL') & (c.variable_2=='estim_total')])\ndisplay(c[(c.variable_1=='estim_por_alumno') & (c.variable_2=='pct_completas')])\ndisplay(tabla('pme_correlaciones_estratificadas'))\ndisplay(Image(filename=str(ROOT/'informes/graficos/06_pme_correlaciones.png')))\ndisplay(Image(filename=str(ROOT/'informes/graficos/07_pme_relaciones.png')))"),
      md('## 5. La Cisterna: fracción completa versus avance\n\nÓscar Encalada tiene cero acciones completas, pero todas están entre 75% y 99%. Se muestran límites de avance; no se imputa un punto medio. La matrícula de 2024 sólo se usa para las estimaciones de 2024.'),
      code("pilotos = tabla('pme_panel_pilotos_2024')\ndisplay(pilotos[pilotos.COD_DEPE.isin([1,2,6])][['RBD','NOM_RBD','MAT_TOTAL','estim_total','estim_por_alumno','pct_completas','avance_min','avance_max']])\ndisplay(Image(filename=str(ROOT/'informes/graficos/09_pme_rangos_la_cisterna.png')))"),
      md('## 6. Compras y observaciones de Melipeuco\n\nNo es una muestra aleatoria nacional: es el universo que definió la auditoría. Los valores son pesos históricos. El histograma por lote y el ponderado por unidades responden a preguntas distintas.'),
      code("compras=tabla('cgr_compras_lotes')\nno_habidos=tabla('cgr_no_habidos')\nprint('Partidas:',len(compras),'Unidades:',compras.cantidad.sum(),'Monto sumado:',compras.valor_clp.sum())\ndisplay(tabla('cgr_resumen_anio'))\ndisplay(tabla('cgr_fuentes_financiamiento'))\ndisplay(Image(filename=str(ROOT/'informes/graficos/01_precios_lotes_unidades.png')))"),
      code("display(no_habidos.groupby('observacion').agg(equipos=('id_fila','size'),valor_clp=('valor_clp','sum')))\nprint('Tasa física:',len(no_habidos)/776,'Tasa monetaria:',no_habidos.valor_clp.sum()/176164029)\ndisplay(Image(filename=str(ROOT/'informes/graficos/04_conciliacion_caso.png')))\ndisplay(tabla('cgr_resumen_modelo'))\ndisplay(Image(filename=str(ROOT/'informes/graficos/02_valor_por_modelo.png')))"),
      md('## 7. Cruce de evidencia y asignación incierta\n\nEl cruce por factura y serie encuentra 476 registros con comodato entre los no habidos. Los 18 equipos sin distribución por escuela no se asignan al DEM por mera igualdad de cantidad. Los intervalos por escuela son límites marginales, no intervalos de confianza.'),
      code("cruce=tabla('cgr_interseccion_comodato_no_habidos')\nprint('Coincidencias:',len(cruce),'Valor:',cruce.valor_clp.sum(),'Proporción del anexo 5:',len(cruce)/600)\ndisplay(tabla('cgr_limites_por_dependencia'))\ndisplay(Image(filename=str(ROOT/'informes/graficos/03_limites_por_destino.png')))"),
      md('## 8. Controles y conclusión\n\nLa diferencia de $1 se conserva. La duplicidad TIC de 61 entregas está incluida en el total observado y no se suma. Se requiere seguimiento para saber qué se recuperó, qué bajas fueron acreditadas y cuál es la pérdida definitiva. Los montos PME y SIMCE no pueden reemplazar esa prueba. Ver informe y solicitudes documentales para descargos y fuentes.'),
      code("import runpy\nrunpy.run_path(str(ROOT/'scripts/verificar.py'),run_name='__main__')\ndisplay(tabla('verificaciones'))"),
      md('## 9. ¿Estuvo bien invertido?\n\nLa tasa de bienes no habidos mide custodia histórica, no eficacia ni pérdida definitiva. Las cuatro categorías de gasto se asignarán a partidas pagadas con documentos de legalidad, objetivo previo, entrega/uso y costo comparable; los casos sin esa evidencia quedan pendientes. Melipeuco presenta una señal fiscalizada de potencial irregularidad en custodia, con seguimiento aún desconocido. SIMCE posterior no demuestra por sí solo impacto de la compra. El marco reproducible está en informes/Marco_evaluacion_gasto.md.')
    ]
    nb=nbformat.v4.new_notebook(cells=cells,metadata={'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},'language_info':{'name':'python','version':sys.version.split()[0]}})
    old=os.getcwd();os.chdir(ROOT)
    shell=InteractiveShell.instance()
    try:
        count=0
        for cell in nb.cells:
            if cell.cell_type!='code':continue
            count+=1
            with capture_output(stdout=True,stderr=True,display=True) as cap:
                result=shell.run_cell(cell.source,store_history=False)
            if result.error_before_exec or result.error_in_exec:
                raise RuntimeError(f'Falla en celda {count}: {result.error_before_exec or result.error_in_exec}\n{cap.stderr}\n{cap.stdout}')
            cell.execution_count=count
            outputs=[]
            if cap.stdout:outputs.append(nbformat.v4.new_output('stream',name='stdout',text=cap.stdout))
            if cap.stderr:outputs.append(nbformat.v4.new_output('stream',name='stderr',text=cap.stderr))
            for output in cap.outputs:
                outputs.append(nbformat.v4.new_output('display_data',data=output.data,metadata=output.metadata))
            cell.outputs=outputs
        nbformat.validate(nb)
        nbformat.write(nb,ROOT/'notebooks/EDA_gasto_educativo.ipynb')
        print(f'Notebook ejecutado: {count} celdas de código, sin errores.')
    finally:os.chdir(old)

if __name__=='__main__':
    notebook()
    html_report()
    workbook()
    print('HTML autónomo y Excel generados.')
