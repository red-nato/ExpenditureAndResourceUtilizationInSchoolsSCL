# %% [markdown]
# # La Cisterna: EDA y revisión documental de acciones PME 2024
#
# **Grupo 03 · Javier Alcaíno, Lucas Riquelme y Renato Varela · 28 de septiembre de 2026**
#
# Este es el **único notebook vigente del análisis**. Reúne pregunta, fuentes, inspección, limpieza con antes/después, panel, EDA y respuesta. Las salidas se obtienen ejecutando las celdas; no son resultados escritos manualmente. W1 se conserva como antecedente fechado.
#
# **Pregunta propuesta:** ¿Cómo se distribuye la proporción de acciones PME 2024 con estimación de $0 o sin implementación declarada completa entre los establecimientos y dimensiones de La Cisterna, y qué casos conviene priorizar para solicitar respaldo documental?
#
# **Por qué cambia.** La pregunta anterior pedía comprobar el uso de recursos y los gastos observados. Faltan rendiciones desagregadas, pagos, recepción y expedientes con materia de cargos. Sí tenemos acciones PME, estimaciones y estados declarados. Adoptamos la **línea B de W1** como análisis principal; la línea A queda como siguiente fase documental. OC y PAS contextualizan la trazabilidad. SIMCE/IDPS describen resultados educativos, sin determinar por sí solos la prioridad documental.
#
# **Decisión y usuarios:** equipos directivos, sostenedores y equipo investigador pueden ordenar solicitudes de respaldo del PME. Se pide explicar la estimación y acreditar la implementación. El indicador no clasifica corrupción, malgasto ni calidad escolar.
#
# **Alcance:** PME 2024 para RBD con acciones observadas; Directorio, PAS y compras 2022–2025 como contexto; SIMCE/IDPS 2023–2025 por aplicación. Importes en CLP nominales. Un año PME no permite calcular variaciones interanuales.
#
# ## Ruta del cuaderno
#
# 0. Metodología: diagrama del análisis y decisiones (empieza aquí).
# 1. Fuentes y vocabulario: qué es cada archivo.
# 2. Load y Profile: procedencia, unidades, llaves y datos originales.
# 3. Explore: diagnóstico previo a limpiar.
# 4. Clean e Impute: reglas, ejemplos y balances antes/después.
# 5. Transform y Validate: KPI y cruces sin multiplicar filas.
# 6. Explore again: distribuciones, faltantes, extremos y asociaciones.
# 7. Respuesta y decisión: segmentos, umbrales, casos y límites.
# 8. Anexos: siete pasos, Five Cs, bitácora y pauta del curso.
#
# Cada sección conecta **pregunta, cálculo e interpretación**. Las correlaciones son diagnósticas y descriptivas, sin inferencia causal.

# %% [markdown]
# ## Metodología: la lógica del análisis (léela antes de las celdas)
#
# **Idea central.** No partimos de «hacer gráficos»: partimos de una **decisión** (¿qué respaldo del PME pedir primero?) y de un **indicador** (KPI) que la sostiene. Después seguimos el ciclo que exige el curso —*Load → Profile → Explore → Clean → Impute → Transform → Validate → Explore again*— porque **el EDA que se interpreta debe hacerse sobre la versión más válida de los datos**, no sobre los archivos tal como llegaron.
#
# La figura resume el recorrido: cada caja dice **qué se hizo** y **por qué**. La flecha punteada indica que la limpieza es iterativa: si el EDA final revela un problema residual, se vuelve a limpiar y a validar.

# %%
import textwrap
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from IPython.display import display
_root = next(p for p in [Path.cwd(), *Path.cwd().parents] if (p/'data/extraidos/la_cisterna').is_dir())
_fig_dir = _root/'analysis/figures'; _fig_dir.mkdir(parents=True, exist_ok=True)
steps = [
    ('1 · Load', 'Cargar 6 fuentes públicas; CSV leídos como texto; huellas SHA-256 verificadas.',
     'Para que el «antes» no desaparezca por una conversión automática.'),
    ('2 · Profile', 'Definir unidad de observación, llaves y tipos de cada fuente.',
     'La unidad decide qué se puede calcular: acción ≠ RBD–año ≠ orden de compra.'),
    ('3 · Explore (1.ª vez)', 'Mirar colas de importes PME y vacíos ocultos (texto vacío ≠ NaN).',
     'Diagnosticar antes de limpiar: orienta las reglas, no saca conclusiones.'),
    ('4 · Clean / Impute', 'Conversión explícita, OC una vez por código, 0 filas eliminadas, sin imputar.',
     'Ausente ≠ cero: imputar fabricaría evidencia de inactividad o bajo desempeño.'),
    ('5 · Transform', 'KPI por acción → colegio → dimensión; agregar cada fuente antes de unir por RBD–año.',
     'Unir acciones con resultados fila a fila multiplicaría filas y montos.'),
    ('6 · Validate', '18 controles automáticos: llaves únicas, filas e importes conservados, KPI entre 0 y 100.',
     'Si una regla se rompe el notebook se detiene: no se interpreta lo no verificado.'),
    ('7 · Explore again', 'Histogramas, densidades, ECDF, IQR/MAD, faltantes, extremos, correlaciones y segmentos.',
     'El EDA que se interpreta usa la versión más válida de los datos.'),
    ('8 · Respuesta', 'Distribución del KPI, umbrales 25/50/75 %, lista de respaldo a solicitar.',
     'Convierte el hallazgo en una decisión de revisión, no en una acusación.'),
]
colors = ['#176B87', '#176B87', '#D28E32', '#488F72', '#488F72', '#9963A5', '#D28E32', '#933D48']

fig, ax = plt.subplots(figsize=(16, 10.2))
ax.set_xlim(0, 16); ax.set_ylim(0, 10.2); ax.axis('off')

# banda superior: decisión y KPI
ax.add_patch(FancyBboxPatch((0.3, 8.75), 15.4, 1.25, boxstyle='round,pad=0.02,rounding_size=0.15',
                            fc='#EAF2F5', ec='#12495C', lw=1.6))
ax.text(8, 9.72, 'PUNTO DE PARTIDA: la decisión, no el gráfico', ha='center', va='center',
        fontsize=12, fontweight='bold', color='#12495C')
ax.text(8, 9.22, '¿Qué respaldo del PME 2024 conviene pedir primero a cada colegio de La Cisterna?',
        ha='center', va='center', fontsize=11)
ax.text(8, 8.93, 'KPI = 100 × acciones con estimación $0 o sin implementación completa (sin doble conteo) / acciones observadas',
        ha='center', va='center', fontsize=10, style='italic', color='#203248')

W, H = 3.6, 2.95
xs = [0.3 + i * 4.0 for i in range(4)]
rows_y = [5.0, 0.8]
pos = []
for i, (title, what, why) in enumerate(steps):
    r, c = divmod(i, 4)
    x, y = xs[c], rows_y[r]
    pos.append((x, y))
    ax.add_patch(FancyBboxPatch((x, y), W, H, boxstyle='round,pad=0.02,rounding_size=0.12',
                                fc='white', ec=colors[i], lw=2.2))
    ax.add_patch(FancyBboxPatch((x, y + H - 0.62), W, 0.62, boxstyle='round,pad=0.02,rounding_size=0.12',
                                fc=colors[i], ec=colors[i], lw=2.2))
    ax.text(x + W / 2, y + H - 0.31, title, ha='center', va='center', color='white', fontsize=12, fontweight='bold')
    ax.text(x + 0.15, y + H - 0.8, 'Qué se hizo', fontsize=9, fontweight='bold', color=colors[i], va='top')
    ax.text(x + 0.15, y + H - 1.08, '\n'.join(textwrap.wrap(what, 38)), fontsize=9.3, va='top', color='#203248', linespacing=1.3)
    ax.text(x + 0.15, y + 1.10, 'Por qué', fontsize=9, fontweight='bold', color=colors[i], va='top')
    ax.text(x + 0.15, y + 0.82, '\n'.join(textwrap.wrap(why, 38)), fontsize=9.3, va='top', color='#203248', linespacing=1.3)

arrow = dict(arrowstyle='-|>', color='#476579', lw=2, mutation_scale=16)
# flechas horizontales dentro de cada fila
for r in range(2):
    for c in range(3):
        x, y = pos[r * 4 + c]
        ax.annotate('', xy=(x + W + 0.38, y + H / 2), xytext=(x + W + 0.02, y + H / 2), arrowprops=arrow)
# 4 -> 5 (codo por la separación entre filas)
x4, y4 = pos[3]; x5, y5 = pos[4]
gy = 4.72
ax.plot([x4 + W / 2 - 0.9, x4 + W / 2 - 0.9], [y4, gy], color='#476579', lw=2)
ax.plot([x4 + W / 2 - 0.9, x5 + W / 2], [gy, gy], color='#476579', lw=2)
ax.annotate('', xy=(x5 + W / 2, y5 + H + 0.02), xytext=(x5 + W / 2, gy), arrowprops=arrow)
# bucle de retorno 7 -> 4 (iteración)
x7, y7 = pos[6]
ly = 4.1
loop = dict(arrowstyle='-|>', color='#933D48', lw=2, mutation_scale=16, ls='--')
ax.plot([x7 + W / 2, x7 + W / 2], [y7 + H, ly], color='#933D48', lw=2, ls='--')
ax.plot([x7 + W / 2, x4 + W / 2 + 0.9], [ly, ly], color='#933D48', lw=2, ls='--')
ax.annotate('', xy=(x4 + W / 2 + 0.9, y4 - 0.02), xytext=(x4 + W / 2 + 0.9, ly), arrowprops=loop)
ax.text(x7 + W / 2 + 0.25, 4.41, 'Si el EDA final revela un problema residual,\nse vuelve a limpiar y a validar (iterativo)',
        fontsize=8.8, color='#933D48', va='center', fontweight='bold', linespacing=1.25)

ax.text(8, 0.32, 'Principios transversales: no imputar montos ni puntajes · no borrar extremos (se marcan y se mide su efecto) · ausencia ≠ cero\n'
                 'correlación ≠ causa · el KPI prioriza la revisión documental, no mide malgasto',
        ha='center', va='center', fontsize=9.8, color='#12495C', fontweight='bold', linespacing=1.4)
ax.set_title('Metodología del análisis: del problema de decisión a la lista de respaldo documental',
             fontsize=14, fontweight='bold', color='#12495C', pad=10)
fig.savefig(_fig_dir/'00_metodologia.png', dpi=160, bbox_inches='tight')
display(fig)
plt.close(fig)

# %% [markdown]
# ### Decisiones metodológicas y alternativas descartadas
#
# | Decisión tomada | Alternativa descartada | Por qué |
# |---|---|---|
# | Pregunta descriptiva sobre señales PME 2024 (línea B de W1) | Comprobar el uso de recursos con rendiciones y pagos (línea A) | Rendiciones, pagos y recepción no están en los datos; las acciones PME sí |
# | KPI = unión de «estimación $0» o «no completa», sin doble conteo | Usar la cobertura de fuentes como KPI | La cobertura mide disponibilidad, no se conecta con una decisión |
# | Un único corte (2024) | KPI de variación interanual | Solo hay un año PME extraído |
# | Agregar cada fuente antes de unir por RBD–año | Unir acciones con resultados fila a fila | Multiplicaría filas y montos |
# | Conservar ceros y ausencias; no imputar | Rellenar con 0 o con la mediana | Un cero declarado no es una celda vacía; las ausencias son estructurales |
# | Marcar extremos con IQR y medir su efecto | Borrarlos para «limpiar» | Un extremo no es un error: puede ser una acción legítima de gran escala |
# | Mediana, IQR y MAD junto a media y desviación estándar | Solo media y desviación | Los importes son muy asimétricos: la media no representa a la acción típica |
# | Correlación en corte 2024, una fila por RBD, con Pearson, Spearman, Kendall y n por pareja | Correlacionar todo el panel | Repetiría el mismo PME en varios años y trataría filas del mismo colegio como independientes |
# | Contar cada orden de compra una vez por código | Sumar el total de la orden por cada ítem | El total se repite por ítem: sobreestima el monto entre 21 y 33 veces (2022–2024) |
# | Umbrales 25/50/75 % y una definición estricta como sensibilidad | Fijar un único 50 % | El 50 % es un supuesto operativo: hay que ver cuánto cambia la lista de colegios |

# %% [markdown]
# ## 1. Fuentes: qué significa cada cosa
#
# **RBD** identifica un establecimiento. **Sostenedor** es quien lo administra: puede cambiar sin cambiar el RBD. **PME** es Plan de Mejoramiento Educativo. Una acción PME es una actividad de mejora registrada por el colegio. **OC** es orden de compra; **PAS**, proceso administrativo sancionatorio; **SIMCE**, resultados de aprendizaje; **IDPS**, indicadores de desarrollo personal y social.
#
# | Fuente | Qué representa una fila | Archivos y cobertura | Para qué sirve | Límite |
# |---|---|---|---|---|
# | Directorio Mineduc | Establecimiento en un año | CSV nacionales y `universo_la_cisterna_2022_2025.json`, 238 RBD–año | Censo, matrícula, dependencia y oferta | No registra ingresos, subvenciones ni pagos |
# | Implementación PME 2024 | Acción declarada | XLSX nacional y `pme_la_cisterna_2024.csv`, 767 acciones de 46 RBD | Estimación total/SEP, dimensión, avance y KPI principal | No acredita gasto ejecutado ni aprobado |
# | PAS Supereduc | Proceso identificado por PA_ID | ZIP/XLSX 2022–2025 y JSON, 110 procesos | Estado, instancia, antecedentes a solicitar | No contiene materia detallada de cargos. Multa no equivale a malgasto |
# | SIMCE, Agencia de Calidad | RBD–año–grado, con medidas por área | RAR 2023–2025 y JSON, 236 filas | Puntajes y disponibilidad por aplicación | Una fila puede existir sin puntaje; no es un resultado único del colegio |
# | IDPS, Agencia de Calidad | RBD–año–grado–indicador | RAR 2023–2025 y JSON, 936 filas | Describir cada indicador y su disponibilidad | No mezclar grados/indicadores. Cambian códigos en 2025 |
# | ChileCompra | En origen, ítem/cotización de OC; al depurar, una OC | Ocho `.7z`: municipio 2022–2024, SLEP 2025 | Montos de órdenes, mecanismo y atribución escolar | OC no acredita pago; cobertura SLEP 2025 parcial |
# | Ficha Comprador | Organismo en una fecha de consulta | `ficha_comprador_2026-09-27.csv`, dos organismos | Contexto de ventanas 2025–2026 | No es una medida por colegio ni por ejercicio 2022–2025 |
#
# **Dimensiones PME:** Gestión Pedagógica agrupa enseñanza/aprendizaje; Liderazgo, conducción institucional; Convivencia Escolar, convivencia y participación; Gestión de Recursos, recursos humanos, materiales y financieros. Clasifican el propósito de acciones, no cuentas contables.
#
# Mineduc advierte que las estimaciones PME pueden cambiar con la rendición y que publicar planes no supone aprobar sus gastos: [catálogo oficial PME](https://liderazgoeducativo.mineduc.cl/bases-de-datos-pme/) y `data/originales/pme_diccionario_2024.pdf`. Implementación y Planificación son etapas distintas. Que el portal publique otros años no significa que estén extraídos en este estudio.
#
# **Carpetas y formatos:** `originales/` conserva paquetes oficiales y archivos descomprimidos. RAR, ZIP y 7Z son contenedores, no fuentes adicionales. Los PDF `ER_...` o `...diccionario` explican campos. `extraidos/la_cisterna/` es la selección comunal. `procesados/la_cisterna/` contiene resultados de cálculos. `fuentes.json` registra procedencia; los manifiestos guardan huellas SHA-256. Un hash verifica identidad, no verdad de un dato.
#
# **Fuera de este repositorio:** el material previo del equipo (Directorios 2015–2021, PME nacional de años anteriores y los informes de Contraloría sobre Melipeuco) **no se incluye ni se usa**; Melipeuco es otro territorio. Las dos versiones de W1, cómo se reconciliaron y la decisión de pasar de la línea A a la B están documentadas en [PROCESS_LOG.md](../PROCESS_LOG.md). No se atribuye esa decisión al docente.

# %% [markdown]
# ## 2. Load: entorno, inventario y datos de entrada
#
# Ejecutar **Restart Kernel and Run All** desde la raíz o desde `analysis/`. Requiere Python, pandas, numpy, matplotlib, scipy, openpyxl y `7z` para releer los archivos OC. No descarga datos ni envía solicitudes. Se conservan los originales y se exportan tablas a `data/procesados/la_cisterna/eda/` y figuras a `analysis/figures/`.

# %%
from pathlib import Path
import sys, json, re, hashlib, csv
from collections import Counter
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap
from matplotlib.ticker import PercentFormatter
from scipy.stats import gaussian_kde, spearmanr, pearsonr, kendalltau
from IPython.display import display, Markdown
ROOT = next(p for p in [Path.cwd(), *Path.cwd().parents] if (p/'data/extraidos/la_cisterna').is_dir())
sys.path.insert(0,str(ROOT/'analysis'))
INP=ROOT/'data/extraidos/la_cisterna'
PRO=ROOT/'data/procesados/la_cisterna'
OUT=PRO/'eda'
FIG=ROOT/'analysis/figures'
OUT.mkdir(parents=True,exist_ok=True)
FIG.mkdir(parents=True,exist_ok=True)
COLORS=['#176B87','#D28E32','#488F72','#9963A5']
plt.rcParams.update({'figure.figsize':(11,4.8),'figure.dpi':110,'axes.spines.top':False,
                     'axes.spines.right':False,'axes.titleweight':'bold','font.size':10})
pd.set_option('display.max_columns',18)
pd.set_option('display.max_rows',65)
pd.set_option('display.float_format',lambda x:f'{x:,.3f}')
def table(df,name=None):
    if name: df.to_csv(OUT/f'{name}.csv',index=False)
    display(df)
def finish(fig,name):
    fig.tight_layout()
    fig.savefig(FIG/f'{name}.png',dpi=160,bbox_inches='tight')
    display(fig)
    plt.close(fig)
def note(text): display(Markdown(text))
def esnum(v,decimals=0): return f'{v:,.{decimals}f}'.replace(',','X').replace('.',',').replace('X','.')
print('Raíz del proyecto:',ROOT.name)
print('Python:',sys.version.split()[0],'pandas:',pd.__version__,'numpy:',np.__version__)

# %% [markdown]
# ### 2.1 Inventario: una fila por archivo relevante localizado
#
# Incluye solo los archivos de la carpeta `data/` de este paquete, de modo que el inventario sea idéntico en cualquier computador. La tabla resume familias; el CSV exportado conserva **cada ruta, tamaño, contenido y uso**. Las copias no se suman como fuentes independientes.

# %%
def describe_file(p):
    n=p.name.lower(); path=str(p).lower()
    if 'directorio' in path and n.startswith('er_'): return 'Directorio','Diccionario de campos y códigos','Metadatos'
    if 'directorio' in n: return 'Directorio','Censo nacional anual','2022–2025 activo; anteriores históricos'
    if n.startswith('universo_'): return 'Directorio local','Censo filtrado de La Cisterna','Usar censo anual 2022–2025'
    if 'chilecompra' in path and p.suffix=='.7z': return 'ChileCompra','Paquete semestral OC por ítem/cotización','Original activo'
    if 'pas' in n or 'supereduc' in n: return 'Supereduc / PAS','Procesos, diccionario o revisión de materia','Activo; no es gasto rechazado'
    if 'simce' in n or 'idps' in n or 'agencia' in n: return 'Agencia','Resultados o cobertura por aplicación','Activo; no sumar copias'
    if 'pme' in n:
        meaning='Acciones, estimaciones y avance, o tablas derivadas'
        if 'diccionario' in n or n.startswith('er '): meaning='Diccionario de variables'
        return 'Mineduc / PME',meaning,'La Cisterna 2024 activo; nacional/pilotos históricos'
    if 'ficha_comprador' in n: return 'Ficha Comprador','Indicadores del organismo en ventanas recientes','Contexto; fuera del panel histórico'
    if n.startswith('oc_'): return 'ChileCompra derivado','Selección, candidatos, cobertura o auditoría OC','Candidatas fuera del agregado principal'
    if n.startswith('panel_'): return 'Integración','Una fila RBD–año tras agregar fuentes','Vigente solo en procesados/la_cisterna'
    if 'manifiesto' in n or n=='fuentes.json': return 'Procedencia','URL, fecha y/o SHA-256','Metadatos'
    if 'diccionario' in n or 'procedencia' in n: return 'Documentación','Campos, llaves y límites','Metadatos'
    return 'Resultado previo','Resumen o control calculado','Derivado; no es fuente nueva'
inventory=[]
for p in sorted((ROOT/'data').rglob('*')):
    if not p.is_file() or p.name.startswith('.') or p.suffix=='.pyc' or 'eda' in p.parts: continue
    fam,meaning,use=describe_file(p)
    inventory.append({'archivo':str(p.relative_to(ROOT)),'familia':fam,'que_contiene':meaning,'uso':use,'MB':round(p.stat().st_size/1e6,3)})
inventory=pd.DataFrame(inventory)
inventory.to_csv(OUT/'inventario_archivos.csv',index=False)
table(inventory.groupby(['familia','que_contiene','uso']).agg(archivos=('archivo','size'),MB=('MB','sum')).reset_index())
note('Inventario individual: `data/procesados/la_cisterna/eda/inventario_archivos.csv`.')

# %% [markdown]
# ### 2.2 Huellas y carga sin ocultar los valores originales
#
# Leemos CSV como texto y conservamos sus cadenas vacías. Los JSON mantienen sus tipos. Así el «antes» no desaparece por una conversión automática. Verificamos huellas de extractos y de originales recuperados; una diferencia detiene el cuaderno.

# %%
manifest=json.loads((INP/'manifiesto_extraccion.json').read_text())
checks=[]
for item in manifest['extractos']:
    assert hashlib.sha256((INP/item['archivo']).read_bytes()).hexdigest()==item['sha256']
    checks.append({'archivo':item['archivo'],'filas_manifestadas':item['filas'],'SHA256_coincide':True})
for item in json.loads((ROOT/'data/manifiesto_recuperacion_local.json').read_text()):
    assert hashlib.sha256((ROOT/item['archivo']).read_bytes()).hexdigest()==item['sha256']
table(pd.DataFrame(checks),'huellas_verificadas')
raw={}
for name,filename in {'directorio':'universo_la_cisterna_2022_2025.json','pas':'procesos_supereduc_la_cisterna_2022_2025.json',
                       'simce':'simce_la_cisterna_2023_2025.json','idps':'idps_la_cisterna_2023_2025.json'}.items():
    raw[name]=pd.DataFrame(json.loads((INP/filename).read_text()))
raw['pme']=pd.read_csv(INP/'pme_la_cisterna_2024.csv',dtype=str,keep_default_na=False)
for name,df in raw.items():
    note(f'**{name.upper()}: {len(df)} filas × {len(df.columns)} columnas.** Primeras dos filas:')
    display(df.head(2))

# %% [markdown]
# ### 2.3 Selección previa: del Directorio nacional al censo comunal
#
# El extracto no es un original nacional. Repetimos la selección: comuna 13109, funcionamiento, matrícula informada, dependencia 2/3/6 y matrícula positiva. RBD 9860 queda fuera por administración delegada, no por resultados. RBD 9830 está solo en el censo 2022–2023. Fuera del censo no significa error.

# %%
national_directory={}; filter_rows=[]
files={2022:'directorio_2022.csv',2023:'directorio_2023.csv',2024:'directorio_2024.csv',
       2025:'20250926_Directorio_Oficial_EE_2025_20250430_WEB.csv'}
for year,name in files.items():
    df=pd.read_csv(ROOT/'data/originales'/name,sep=';',dtype=str,keep_default_na=False,encoding='utf-8-sig')
    national_directory[year]=df
    stages=[('Nacional',pd.Series(True,index=df.index)),('La Cisterna',df.COD_COM_RBD.eq('13109')),
            ('Funcionamiento y matrícula informada',df.COD_COM_RBD.eq('13109')&df.ESTADO_ESTAB.eq('1')&df.MATRICULA.eq('1'))]
    valid=stages[-1][1]&df.COD_DEPE.isin(['2','3','6'])&pd.to_numeric(df.MAT_TOTAL).gt(0)
    stages.append(('Censo de dependencias seleccionadas',valid))
    for stage,mask in stages: filter_rows.append({'anio':year,'etapa':stage,'filas':int(mask.sum())})
    assert set(df.loc[valid,'RBD'].astype(int))==set(raw['directorio'].loc[raw['directorio'].anio.eq(year),'rbd'])
table(pd.DataFrame(filter_rows),'seleccion_directorio_antes_despues')
table(pd.concat([d.loc[d.RBD.isin(['9830','9860']),['RBD','NOM_RBD','COD_DEPE','ESTADO_ESTAB','MAT_TOTAL']].assign(anio=y)
                 for y,d in national_directory.items()],ignore_index=True),'casos_9830_9860')

# %% [markdown]
# ### 2.4 Cotejo PME con el XLSX oficial recuperado
#
# Leemos el XLSX nacional fila a fila y filtramos el censo 2024. Comprobamos localizador, RBD, importes y estado. Recuperamos nombre y descripción de la acción para revisar extremos. `fila_excel` ubica una fila de esta versión; no es un ID oficial estable entre años.

# %%
from openpyxl import load_workbook
wb=load_workbook(ROOT/'data/originales/pme_implementacion_2024.xlsx',read_only=True,data_only=True)
rows=wb.active.iter_rows(values_only=True); headers=next(rows)
rbd_index=headers.index('RBD')
scope_2024=set(raw['directorio'].loc[raw['directorio'].anio.eq(2024),'rbd'])
official_actions=[]; national_action_count=0
for excel_row,values in enumerate(rows,start=2):
    national_action_count+=1
    if values[rbd_index] in scope_2024: official_actions.append(dict(zip(headers,values))|{'fila_excel':excel_row})
wb.close()
official_pme=pd.DataFrame(official_actions)
assert len(official_pme)==len(raw['pme'])==767
comparison=raw['pme'].assign(fila_excel=lambda d:d.fila_excel.astype(int)).merge(
    official_pme[['fila_excel','RBD','ESTIM_TOTAL','ESTIM_SEP','NIV_IMPLEM']],on='fila_excel',suffixes=('_extracto','_oficial'),validate='one_to_one')
for col in ['RBD','ESTIM_TOTAL','ESTIM_SEP']:
    assert np.allclose(pd.to_numeric(comparison[col+'_extracto']),comparison[col+'_oficial'])
assert comparison.NIV_IMPLEM_extracto.eq(comparison.NIV_IMPLEM_oficial).all()
note(f'**Cotejo:** {national_action_count:,} acciones nacionales; {len(official_pme)} del censo 2024; coincidencia total de campos contrastados.')
display(official_pme[['fila_excel','RBD','DIMENSION','NOM_ACTIVIDAD','NIV_IMPLEM','ESTIM_TOTAL']].head(4))

# %% [markdown]
# ### 2.5 OC: unidad de observación y relectura de los ocho originales
#
# Agrupamos por `codigoOC`, validando que coincidan monto, moneda, procedencia, unidad y estado. Se reúne el texto de todos los ítems. Luego se selecciona EDUCACIÓN municipal o una escuela inequívoca SLEP sin otra comuna mencionada. Las OC compartidas se conservan sin reparto escolar especulativo.
#
# La auditoría muestra filas e importes antes/después **bajo los mismos filtros de selección, estado y moneda**. Repetir el total por ítem no prueba compras duplicadas. Se conservan aparte las candidatas fuera del filtro.

# %%
from scripts.integrar_chilecompra import run as leer_oc_originales
from scripts.pipeline_la_cisterna import run as reproducir_panel_base
leer_oc_originales()
reproducir_panel_base(INP,PRO)
raw['oc']=pd.read_csv(PRO/'oc_escuelas_publicas.csv',dtype=str,keep_default_na=False)
oc_audit=pd.read_csv(PRO/'oc_limpieza_antes_despues.csv')
table(oc_audit,'auditoria_oc_originales')

# %% [markdown]
# ## 3. Profile y Explore: problemas antes de limpiar
#
# ### 3.1 Llaves, tipos y ausencias aparentes
#
# Repetir RBD es esperado en acciones/procesos. Una llave depende de lo que representa la fila. `.isna()` no reconoce una cadena vacía como ausente; contamos ambas representaciones. `medidas_originales` de SIMCE debe desplegarse antes de diagnosticar sus puntajes.

# %%
keys={'directorio':['anio','rbd'],'pas':['pa_id'],'pme':['fila_excel'],
      'simce':['anio','rbd','grado'],'idps':['anio','rbd','grado','indicador_codigo_original'],'oc':['codigo_oc']}
profile_rows=[]
for name,df in raw.items():
    for col in df:
        v=df[col]; blank=v.map(lambda x:isinstance(x,str) and not x.strip())
        profile_rows.append({'fuente':name,'campo':col,'tipo_antes':str(v.dtype),'filas':len(df),
                             'nulos_reales':int(v.isna().sum()),'texto_vacio':int(blank.sum()),'distintos':len(set(v.map(repr)))})
profile=pd.DataFrame(profile_rows)
table(pd.DataFrame([{'fuente':name,'filas':len(df),'llave':' + '.join(keys[name]),
                     'llaves_repetidas_adicionales':int(df.duplicated(keys[name]).sum())} for name,df in raw.items()]),'llaves_antes')
table(profile.loc[profile.texto_vacio.gt(0)|profile.nulos_reales.gt(0)|profile.campo.isin(['RBD','ESTIM_TOTAL','NIV_IMPLEM','matricula_total'])],'perfil_campos_seleccion')
profile.to_csv(OUT/'perfil_todos_los_campos_antes.csv',index=False)

# %%
simce_records=[]
for row in raw['simce'].to_dict('records'):
    for key,value in row['medidas_originales'].items():
        if not key.startswith('prom_'): continue
        area=re.match(r'prom_(lect|mate|hist)',key).group(1)
        suffix=key.removeprefix('prom_')
        simce_records.append({k:v for k,v in row.items() if k!='medidas_originales'}|
                             {'area':area,'puntaje_original':value,'campo_original':key,
                              'alumnos_original':row['medidas_originales'].get('nalu_'+suffix),
                              'marca_original':row['medidas_originales'].get('marca_'+suffix)})
simce_raw_long=pd.DataFrame(simce_records)
table(pd.DataFrame([{'medida':name,'filas':len(s),'isna_sin_limpiar':int(s.isna().sum()),'vacios_texto':int(s.eq('').sum())}
                    for name,s in [('IDPS',raw['idps'].promedio_original),('SIMCE por área',simce_raw_long.puntaje_original)]]),'faltantes_aparentes_antes')
table(raw['pme'].NIV_IMPLEM.value_counts(dropna=False).rename_axis('estado_original').reset_index(name='acciones'))
table(raw['oc'].groupby(['moneda_oc','estado_oc']).size().reset_index(name='ordenes'))

# %% [markdown]
# ### 3.2 Exploración provisional: asimetría y vacíos escondidos
#
# La primera figura muestra la cola de importes PME antes de tratamientos. La segunda contrasta ausencia detectada por `.isna()` con texto vacío. Esto orienta reglas de limpieza, sin sacar aún conclusiones de asociaciones entre fuentes.

# %%
fig,axes=plt.subplots(1,2,figsize=(12,4.2))
axes[0].hist(pd.to_numeric(raw['pme'].ESTIM_TOTAL)/1e6,bins=35,color=COLORS[0])
axes[0].set(title='PME: importes antes de tratamientos',xlabel='Estimación por acción (millones CLP)',ylabel='Acciones')
bm=pd.DataFrame({'Solo isna':[raw['idps'].promedio_original.isna().sum(),simce_raw_long.puntaje_original.isna().sum()],
                 'Texto vacío':[raw['idps'].promedio_original.eq('').sum(),simce_raw_long.puntaje_original.eq('').sum()]},index=['IDPS','SIMCE por área'])
bm.plot.bar(ax=axes[1],color=COLORS[:2],rot=0)
axes[1].set(title='Ausencia guardada como texto',ylabel='Celdas sin resultado',xlabel='')
finish(fig,'01_diagnostico_antes')

# %% [markdown]
# ## 4. Clean e Impute: reglas justificadas
#
# | Problema | Tratamiento | Razón y conservación |
# |---|---|---|
# | Números como texto | Conversión explícita; errores detienen el análisis | RBD/año enteros; originales intactos |
# | Vacío, espacios o None | Normalizar a NaN | No confundir ausencia y cero; conservar valor original |
# | IDPS cambia códigos en 2025 | AM/1, CC/2, PF/3 y HV/4 se homologan por nombre | Conservar código y versión originales |
# | Total OC repetido | Una vez por código, validando atributos | No sumar el total por cada ítem |
# | OC sin RBD único | Mantener en agregado de la unidad, fuera del monto escolar | No inventar destinatarios |
# | Importe cero | Conservar | Declaración cero no equivale a celda vacía |
# | Extremo | Marcar y revisar; sensibilidad separada | No borrar una observación solo por ser grande |
# | Importe/puntaje ausente | No imputar | No hay fundamento para inventar un valor |
#
# Un conteo de filas sin coincidencia puede ser cero dentro de la cobertura. Un monto desconocido permanece ausente. Tener OC vinculadas pero ninguna elegible en CLP **no permite imputar monto cero**.

# %%
conversion_log=[]
def clean_numeric(s,source,field,integer=False):
    norm=s.map(lambda v:np.nan if v is None or (isinstance(v,str) and not v.strip()) else v)
    parsed=pd.to_numeric(norm,errors='coerce')
    failed=norm.notna()&parsed.isna()
    conversion_log.append({'fuente':source,'campo':field,'tipo_antes':str(s.dtype),
                           'tipo_despues':'Int64' if integer else str(parsed.dtype),
                           'nulos_antes_isna':int(s.isna().sum()),
                           'vacios_antes':int(s.map(lambda v:isinstance(v,str) and not v.strip()).sum()),
                           'nulos_despues':int(parsed.isna().sum()),'fallos_conversion':int(failed.sum())})
    assert not failed.any(),f'{source}.{field}: tokens inesperados {norm[failed].unique()}'
    if integer:
        assert parsed.dropna().mod(1).eq(0).all()
        return parsed.astype('Int64')
    return parsed
directory=raw['directorio'].copy(deep=True); pas=raw['pas'].copy(deep=True); pme=raw['pme'].copy(deep=True)
simce=simce_raw_long.copy(deep=True); idps=raw['idps'].copy(deep=True); oc=raw['oc'].copy(deep=True)
for name,df,cols in [('directorio',directory,['anio','rbd','matricula_total']),('pas',pas,['archivo_anio','pa_id','rbd','anio_ingreso','anio_termino']),
                     ('pme',pme,['RBD','fila_excel','COD_DEPE']),('simce',simce,['anio','rbd']),('idps',idps,['anio','rbd']),('oc',oc,['anio','rbd_asignado','trato_directo'])]:
    for col in cols: df[col]=clean_numeric(df[col],name,col,integer=True)
pme['estimacion_clp']=clean_numeric(pme.ESTIM_TOTAL,'pme','ESTIM_TOTAL')
pme['sep_clp']=clean_numeric(pme.ESTIM_SEP,'pme','ESTIM_SEP')
simce['puntaje']=clean_numeric(simce.puntaje_original,'simce','puntaje_original')
simce['alumnos']=clean_numeric(simce.alumnos_original,'simce','alumnos_original',integer=True)
idps['puntaje']=clean_numeric(idps.promedio_original,'idps','promedio_original')
for col in ['monto_total_oc_clp','monto_oc_vigente_clp']: oc[col]=clean_numeric(oc[col],'oc',col)
for df in [simce,idps]: df['grado']=df.grado.str.strip().str.lower()
pme['dimension']=pme.DIMENSION.str.strip()
pme['estado_declarado']=pme.NIV_IMPLEM.str.strip().replace('',pd.NA)
idps['indicador_std']=idps.indicador_codigo_original.astype(str).map({'AM':'AM','1':'AM','CC':'CC','2':'CC','PF':'PF','3':'PF','HV':'HV','4':'HV'})
assert idps.indicador_std.notna().all()
assert idps.groupby('indicador_std').indicador.nunique().eq(1).all()
directory['sector']=np.where(directory.dependencia.eq('particular_subvencionado'),'Particular subvencionado','Público')
pme=pme.merge(directory.loc[directory.anio.eq(2024),['rbd','nombre','sector','oferta','matricula_total']],left_on='RBD',right_on='rbd',how='left',validate='many_to_one')
pme=pme.merge(official_pme[['fila_excel','NOM_ACTIVIDAD','DESC_ACTIVIDAD','CONVENIO_SEP']],on='fila_excel',how='left',validate='one_to_one')
table(pd.DataFrame(conversion_log),'limpieza_por_campo')

# %% [markdown]
# ### 4.1 Antes/después: filas, ejemplos e importes
#
# No se eliminan acciones ni puntajes: los problemas de los extractos son representación, tipo y granularidad. PME conserva montos. En OC se mide la diferencia real de sumar por fila en vez de por orden. No se simulan errores inexistentes para presentar mejoras artificiales.

# %%
clean_frames={'directorio':directory,'pas':pas,'pme':pme,'simce':simce,'idps':idps,'oc':oc}
table(pd.DataFrame([{'fuente':name,'filas_antes_misma_unidad':len(simce_raw_long if name=='simce' else raw[name]),
                    'filas_despues':len(df),'filas_eliminadas':len(simce_raw_long if name=='simce' else raw[name])-len(df)}
                   for name,df in clean_frames.items()]),'limpieza_balance_filas')
table(pd.DataFrame([
    ['PME','RBD',repr(raw['pme'].RBD.iloc[0]),repr(int(pme.RBD.iloc[0])),'Texto a entero'],
    ['PME','ESTIM_TOTAL',repr(raw['pme'].ESTIM_TOTAL.iloc[0]),repr(int(pme.estimacion_clp.iloc[0])),'Cero conservado'],
    ['IDPS','promedio_original',repr(''),'NaN','Vacío no es cero'],
    ['SIMCE','puntaje_original',repr(''),'NaN','Vacío no es cero'],
    ['IDPS 2025','indicador','1','AM','Homologación por nombre']
],columns=['fuente','campo','antes','despues','regla']),'ejemplos_antes_despues')
oc_audit['factor_sobresuma']=oc_audit.suma_incorrecta_por_fila_clp/oc_audit.suma_correcta_por_oc_clp
table(oc_audit[['anio','filas_originales_seleccionadas','oc_unicas_seleccionadas','suma_incorrecta_por_fila_clp','suma_correcta_por_oc_clp','factor_sobresuma']],'oc_efecto_limpieza')
fig,axes=plt.subplots(1,2,figsize=(12,4.4)); x=np.arange(len(oc_audit))
axes[0].bar(x-.18,oc_audit.suma_incorrecta_por_fila_clp/1e6,.36,label='Total repetido por fila',color=COLORS[1])
axes[0].bar(x+.18,oc_audit.suma_correcta_por_oc_clp/1e6,.36,label='Una vez por OC',color=COLORS[0])
axes[0].set(xticks=x,xticklabels=oc_audit.anio,title='OC: efecto de corregir granularidad',ylabel='Millones CLP nominales'); axes[0].legend(fontsize=8)
axes[1].hist(pd.to_numeric(raw['pme'].ESTIM_TOTAL)/1e6,bins=30,histtype='step',lw=3,label='Antes')
axes[1].hist(pme.estimacion_clp/1e6,bins=30,histtype='step',lw=1.5,ls='--',label='Después')
axes[1].set(title='PME: importes preservados, curvas coinciden',xlabel='Estimación por acción (millones CLP)',ylabel='Acciones'); axes[1].legend()
finish(fig,'02_limpieza_antes_despues')

# %% [markdown]
# ### 4.2 Calidad residual: reglas de dominio y fechas PAS
#
# Comprobamos negativos, SEP mayor que total, IDPS fuera de 0–100 y fechas. SIMCE no comparte la escala IDPS: no se impone un techo 100. En PAS el año del archivo puede diferir del término; eso no es un error automático. La materia de cargos no se inventa a partir de actividad/programa.

# %%
quality=[('PME importes negativos',int(pme.estimacion_clp.lt(0).sum())),('PME SEP mayor que total',int(pme.sep_clp.gt(pme.estimacion_clp).sum())),
         ('PME duplicado exacto marcado en extracción',int(raw['pme'].duplicado_exacto_adicional.eq('True').sum())),
         ('PME monto o estado ausente',int((pme.estimacion_clp.isna()|pme.estado_declarado.isna()).sum())),
         ('IDPS fuera de 0–100',int((idps.puntaje.notna()&~idps.puntaje.between(0,100)).sum())),
         ('OC CLP negativos',int(oc.monto_total_oc_clp.lt(0).sum())),
         ('SIMCE puntaje sin alumnos positivos',int((simce.puntaje.notna()&~simce.alumnos.gt(0).fillna(False)).sum()))]
table(pd.DataFrame(quality,columns=['control','casos']),'reglas_dominio')
assert all(n==0 for _,n in quality[:2]+quality[3:6])
table(pas.groupby(['archivo_anio','anio_ingreso','anio_termino'],dropna=False).size().reset_index(name='procesos'),'pas_anios_distintos')

# %% [markdown]
# ## 5. Transform: KPI, unidades y panel
#
# ### 5.1 Definición del KPI que sí podemos defender
#
# Para cada acción, definimos **Z = estimación total igual a cero** y **N = no se declara implementación completa**. La señal es **Z o N**. Se cuenta una sola vez si cumple ambas.
#
# **KPI = 100 × acciones con Z o N / total de acciones observadas del colegio o dimensión en 2024.** Se reportan numerador, denominador, componentes y superposición. Un colegio sin acciones en el extracto no tiene KPI calculable. «No completa» incluye niveles parciales/avanzados, no significa «sin ejecución». Si hubiese estado ausente se mostraría separado como no declarado; monto ausente no se convierte a cero.
#
# **Por qué este KPI:** señala acciones cuyo registro requiere explicación o respaldo, vinculando el resultado a una decisión concreta. Es más útil para este propósito que medir solo presencia de una fuente. La prioridad no representa probabilidad de infracción.
#
# **Regla propuesta, no validada:** revisar primero colegios con KPI ≥ 50%; dentro del grupo, ordenar por número de acciones señaladas, conservando siempre el denominador. Contrastar también umbrales 25% y 75%. El 50% es un supuesto operativo, no una norma ni un corte aprendido de resultados. Los colegios sin PME siguen una cola distinta de verificación de cobertura/aplicabilidad, no se clasifican con KPI cero.

# %%
pme['completa']=pme.estado_declarado.eq('Implementación completa: 100%').fillna(False)
pme['sin_estado']=pme.estado_declarado.isna()
pme['no_completa']=~pme.completa
pme['estimacion_cero']=pme.estimacion_clp.eq(0)
pme['ambas']=pme.no_completa & pme.estimacion_cero
pme['senal_revision']=pme.no_completa | pme.estimacion_cero
pme['log10_estimacion_mas1']=np.log10(pme.estimacion_clp+1)

def aggregate_pme(group_cols):
    g=pme.groupby(group_cols,dropna=False).agg(
        acciones=('fila_excel','size'),completas=('completa','sum'),no_completas=('no_completa','sum'),
        estimacion_cero=('estimacion_cero','sum'),ambas=('ambas','sum'),senal_revision=('senal_revision','sum'),
        sin_estado=('sin_estado','sum'),estimacion_mediana_clp=('estimacion_clp','median'),
        estimacion_total_clp=('estimacion_clp',lambda s:s.sum(min_count=1)),rbd_observados=('RBD','nunique')).reset_index()
    for c in ['completas','no_completas','estimacion_cero','senal_revision']:
        g['pct_'+c]=100*g[c]/g.acciones
    return g
dimensions=aggregate_pme(['dimension'])
schools=aggregate_pme(['rbd','nombre','sector','oferta'])
school_dimensions=aggregate_pme(['rbd','nombre','sector','oferta','dimension'])
table(pd.crosstab(pme.estimacion_cero,pme.no_completa).rename_axis(index='Estimación cero',columns='No declarada completa').reset_index(),'superposicion_senales')
note(f'**Unión sin doble conteo:** {pme.no_completa.sum()} no completas + {pme.estimacion_cero.sum()} con $0 '
     f'− {pme.ambas.sum()} con ambas = **{pme.senal_revision.sum()}/{len(pme)} ({100*pme.senal_revision.mean():.1f}%)**. '
     f'Hay {pme.sin_estado.sum()} estados ausentes.')
table(dimensions,'kpi_por_dimension')
schools.to_csv(OUT/'kpi_por_colegio.csv',index=False)
school_dimensions.to_csv(OUT/'kpi_por_colegio_dimension.csv',index=False)

# %% [markdown]
# ### 5.2 Cruces: agregar antes de unir y preservar grado/indicador
#
# El panel base `panel_rbd_anual.csv` tiene una fila RBD–año. No unimos directamente 767 acciones con 936 filas IDPS: eso multiplicaría montos y conteos. PME se agrega primero por colegio y se une **solo a 2024**. Los puntajes se mantienen separados por grado y área/indicador. No se calcula un promedio SIMCE entre grados ni un IDPS general.
#
# Se crea `panel_rbd_anual_eda.csv`, extensión analítica del panel base, con KPI y columnas separadas de resultados. El panel base sigue siendo trazable al pipeline existente. Matrícula es la total del establecimiento, no el número evaluado en un grado; no la usaremos como ponderador de puntajes.

# %%
panel_base=pd.read_csv(PRO/'panel_rbd_anual.csv')
panel=panel_base.merge(directory[['anio','rbd','sector']],on=['anio','rbd'],how='left',validate='one_to_one')
kpi_panel=schools[['rbd','acciones','senal_revision','pct_senal_revision','pct_completas','pct_estimacion_cero','estimacion_mediana_clp']].assign(anio=2024)
kpi_panel=kpi_panel.rename(columns={c:'pme_'+c for c in kpi_panel if c not in ['anio','rbd']})
panel=panel.merge(kpi_panel,on=['anio','rbd'],how='left',validate='one_to_one')
simce_wide=simce.pivot(index=['anio','rbd'],columns=['grado','area'],values='puntaje')
simce_wide.columns=['simce_'+g+'_'+a for g,a in simce_wide.columns]
idps_wide=idps.pivot(index=['anio','rbd'],columns=['grado','indicador_std'],values='puntaje')
idps_wide.columns=['idps_'+g+'_'+i for g,i in idps_wide.columns]
panel=panel.merge(simce_wide.reset_index(),on=['anio','rbd'],how='left',validate='one_to_one')
panel=panel.merge(idps_wide.reset_index(),on=['anio','rbd'],how='left',validate='one_to_one')
panel['pme_log10_total_mas1']=np.log10(panel.pme_estimacion_clp_2024+1)
panel['pme_estimacion_por_matricula_clp']=panel.pme_estimacion_clp_2024/panel.matricula_total
panel['oc_log10_monto_mas1']=np.log10(panel.oc_monto_vigente_clp_rbd_vinculado+1)
table(panel.head(6))
table(pd.DataFrame({'tabla':['Directorio','PME acciones','PME colegios','SIMCE por área','IDPS','Panel'],
                    'filas':[len(directory),len(pme),len(schools),len(simce),len(idps),len(panel)],
                    'unidad':['RBD–año','Acción 2024','RBD 2024','RBD–año–grado–área','RBD–año–grado–indicador','RBD–año']}),'unidades_y_joins')

# %% [markdown]
# ## 6. Validate: pruebas antes de interpretar
#
# Comprobamos unicidad, pertenencia al censo del año, conservación de filas/importes, intersección de señales, rangos del KPI y ausencia de PME fuera de 2024. Reconciliamos las OC vinculadas con el panel usando tolerancia de un centavo. Las reglas fallan explícitamente si se rompe una condición.
#
# Releemos también PAS nacionales para comprobar que la selección no perdió un RBD de 2022–2023 al usar por error el censo 2024. Ninguna coincidencia por RBD se da por segura sin verificar año y unidad.

# %%
pas_original_ids=[]; pas_source_audit=[]
for year in range(2022,2026):
    wb=load_workbook(ROOT/f'data/originales/pas_{year}.xlsx',read_only=True,data_only=True)
    iterator=wb.active.iter_rows(values_only=True); header=next(iterator)
    rbd_i=header.index('EE_RBD'); pa_i=header.index('PA_ID')
    census=set(directory.loc[directory.anio.eq(year),'rbd']); selected=[]; n=0
    for values in iterator:
        n+=1
        if values[rbd_i] in census: selected.append(values[pa_i])
    wb.close()
    pas_original_ids.extend(selected)
    assert set(selected)==set(pas.loc[pas.archivo_anio.eq(year),'pa_id'])
    pas_source_audit.append({'anio_archivo':year,'filas_nacionales':n,'procesos_censo_anual':len(selected),'coincide_extracto':True})
table(pd.DataFrame(pas_source_audit),'pas_cotejo_original')

# %%
validation=[]
def check(label,condition):
    ok=bool(condition); validation.append({'control':label,'resultado':'OK' if ok else 'FALLA'})
    assert ok,label
check('238 filas y llave RBD–año única',len(panel)==238 and not panel.duplicated(['anio','rbd']).any())
check('Censo 60/60/59/59',panel.groupby('anio').size().to_dict()=={2022:60,2023:60,2024:59,2025:59})
check('767 acciones y localizador único',len(pme)==767 and pme.fila_excel.is_unique)
check('PME enlazado con censo 2024',pme.nombre.notna().all())
check('SIMCE llave completa única',not simce.duplicated(['anio','rbd','grado','area']).any())
check('IDPS llave homologada única',not idps.duplicated(['anio','rbd','grado','indicador_std']).any())
check('PAS sin pérdida desde archivos originales',set(pas_original_ids)==set(pas.pa_id) and pas.pa_id.is_unique)
check('OC código único',oc.codigo_oc.is_unique)
check('KPI entre 0 y 100',schools.pct_senal_revision.between(0,100).all())
check('Unión de señales sin doble conteo',pme.senal_revision.sum()==pme.no_completa.sum()+pme.estimacion_cero.sum()-pme.ambas.sum())
check('PME restringido a 2024',panel.loc[panel.anio.ne(2024),'pme_pct_senal_revision'].isna().all())
check('Acciones agregadas conservadas',panel.pme_acciones.sum()==len(pme))
check('Importes PME conservados',np.isclose(panel.pme_estimacion_clp_2024.sum(),pme.estimacion_clp.sum(),rtol=0,atol=.01))
linked=oc.loc[oc.rbd_asignado.notna()]
check('OC vinculadas conservadas',panel.oc_rbd_vinculadas.sum()==len(linked)==637)
check('Importes OC vinculados conservados',np.isclose(panel.oc_monto_vigente_clp_rbd_vinculado.sum(),linked.monto_oc_vigente_clp.sum(),rtol=0,atol=.01))
check('Dimensiones concilian con W1',dimensions.acciones.sum()==767 and dimensions.completas.sum()==469 and dimensions.estimacion_cero.sum()==149)
for name,df in [('simce',simce),('idps',idps)]:
    join=df.merge(directory[['anio','rbd']],on=['anio','rbd'],how='left',indicator=True,validate='many_to_one')
    check(name+' pertenece al censo de cada año',join._merge.eq('both').all())
table(pd.DataFrame(validation),'validaciones')
panel.to_csv(OUT/'panel_rbd_anual_eda.csv',index=False)
pme.to_csv(OUT/'pme_acciones_limpias.csv',index=False)
simce.to_csv(OUT/'simce_por_area_limpio.csv',index=False)
idps.to_csv(OUT/'idps_limpio.csv',index=False)
oc.to_csv(OUT/'oc_limpias.csv',index=False)

# %% [markdown]
# ### 6.1 Validación de los cruces (joins): llave, cardinalidad, cobertura y filas
#
# Cada cruce declara su llave y su cardinalidad esperada (`validate=` de pandas la hace cumplir) y se mide cuántas llaves cruzan, cuántas quedan sin par y cómo cambian las filas. Un cruce sin par no es un error: indica cobertura. «Solo derecha» son registros del censo sin dato en la otra fuente; «solo izquierda» son registros de la fuente sin par en el censo.

# %%
c_keys=set(zip(directory.anio.astype(int),directory.rbd.astype(int)))
c24=set(directory.loc[directory.anio.eq(2024),'rbd'].astype(int))
join_rows=[]
def add_join(cruce,llave,cardinalidad,antes,despues,izq,der,lectura):
    L,R=set(izq),set(der)
    join_rows.append({'cruce':cruce,'llave':llave,'cardinalidad esperada':cardinalidad,'filas antes':int(antes),'filas después':int(despues),
                      'llaves izquierda':len(L),'llaves derecha':len(R),'llaves que cruzan':len(L&R),
                      'solo izquierda':len(L-R),'solo derecha':len(R-L),'lectura':lectura})
add_join('PME (acciones) → Directorio 2024','RBD','muchos a uno',len(raw['pme']),len(pme),pme.RBD.dropna().astype(int),c24,
         'Todas las acciones cruzan; «solo derecha» son colegios del censo sin acciones en el extracto (cobertura, no error)')
for name,df in [('SIMCE',simce),('IDPS',idps)]:
    merged=df.merge(directory[['anio','rbd']],on=['anio','rbd'],how='left',validate='many_to_one')
    add_join(f'{name} → Directorio','RBD + año','muchos a uno',len(df),len(merged),zip(df.anio.astype(int),df.rbd.astype(int)),c_keys,
             'Todas las filas pertenecen al censo del año; «solo derecha» son colegios-año sin resultado publicado (sin fila, no cero)')
add_join('KPI por colegio → panel','RBD + año 2024','uno a uno',len(panel_base),len(panel),zip([2024]*len(schools),schools.rbd.astype(int)),c_keys,
         'El KPI se une solo a 2024; los demás años quedan sin KPI (PME no recopilado)')
add_join('OC con RBD único → panel','RBD + año','muchos a uno (se agrega antes)',len(oc),int(panel.oc_rbd_vinculadas.sum()),
         zip(linked.anio.astype(int),linked.rbd_asignado.astype(int)),c_keys,
         'Solo se asignan las OC con un único RBD; el resto no se reparte (no se inventa destinatario)')
add_join('PAS → panel','RBD + año del archivo','muchos a uno (se agrega antes)',len(pas),int(panel.pas_filas_archivo_anual.sum()),
         zip(pas.archivo_anio.astype(int),pas.rbd.astype(int)),c_keys,
         'Un «solo izquierda» indica un proceso de un RBD que no está en el censo de ese año')
assert all(r['filas antes']==r['filas después'] for r in join_rows[:4]),'un cruce fila a fila cambió el número de filas'
table(pd.DataFrame(join_rows),'joins_validacion')
note('En los cuatro cruces fila a fila (PME, SIMCE, IDPS y KPI→panel) las filas antes y después coinciden: ningún cruce duplica ni pierde registros. '
     'En OC y PAS se compara el total de registros con los efectivamente asignados al panel.')

# %% [markdown]
# ## 7. Explore again: estadística descriptiva robusta
#
# Después de validar, miramos distribuciones completas. **Media y desviación estándar** son sensibles a extremos. **Mediana, IQR (Q75−Q25) y MAD (mediana de desviaciones absolutas)** describen el centro/dispersión con menor sensibilidad. Q90/Q95/Q99 muestran la cola. El número de observaciones válidas y los ceros acompañan cada resumen.
#
# PME se resume por acción y OC por código único. SIMCE e IDPS se resumen dentro de año–grado–área/indicador. No tratamos los 238 RBD–año como 238 escuelas independientes.

# %%
def summarize(s):
    x=pd.to_numeric(s,errors='raise').dropna().astype(float)
    if not len(x): return pd.Series({'n':0,'faltantes':len(s)})
    return pd.Series({'n':len(x),'faltantes':len(s)-len(x),'ceros':int(x.eq(0).sum()),'media':x.mean(),
                      'mediana':x.median(),'desv_est':x.std(),'IQR':x.quantile(.75)-x.quantile(.25),
                      'MAD':(x-x.median()).abs().median(),'min':x.min(),'Q25':x.quantile(.25),
                      'Q75':x.quantile(.75),'Q90':x.quantile(.9),'Q95':x.quantile(.95),'Q99':x.quantile(.99),'max':x.max()})
money_summary=pd.DataFrame({
    'PME estimación por acción 2024':summarize(pme.estimacion_clp),
    'OC nominal CLP (todo estado, 2022–25)':summarize(oc.monto_total_oc_clp),
    'OC CLP aceptada/recepción (2022–25)':summarize(oc.monto_oc_vigente_clp),
    'Matrícula por colegio 2024':summarize(panel.loc[panel.anio.eq(2024),'matricula_total'])}).T.reset_index(names='variable')
table(money_summary,'descriptivos_generales')
def grouped_summary(df,groups,col):
    rows=[]
    for key,g in df.groupby(groups,dropna=False):
        key=key if isinstance(key,tuple) else (key,)
        rows.append(dict(zip(groups,key))|summarize(g[col]).to_dict())
    return pd.DataFrame(rows)
table(grouped_summary(pme,['dimension'],'estimacion_clp'),'descriptivos_pme_dimension')
table(grouped_summary(oc,['anio'],'monto_oc_vigente_clp'),'descriptivos_oc_anio')
table(grouped_summary(simce,['anio','grado','area'],'puntaje'),'descriptivos_simce')
idps_stats=grouped_summary(idps,['anio','grado','indicador_std'],'puntaje')
idps_stats.to_csv(OUT/'descriptivos_idps.csv',index=False)
table(idps_stats.loc[idps_stats.anio.eq(2024)])
note(f'**Lectura PME:** la media es ${esnum(pme.estimacion_clp.mean())} y la mediana ${esnum(pme.estimacion_clp.median())}. '
     'La diferencia muestra asimetría; el promedio solo no representa bien una acción típica. '
     'Los resúmenes globales OC son mezcla de años y coberturas: se usan para diagnosticar la escala, no para inferir crecimiento del gasto.')

# %% [markdown]
# ### 7.1 Histogramas, densidades y ECDF de PME
#
# El histograma lineal conserva ceros y extremos. La densidad se calcula sobre **log10 de importes positivos**, separando la masa de ceros: el suavizado no modela bien una masa puntual en cero. Densidad es forma relativa, no número de acciones. La ECDF muestra qué fracción queda bajo cada importe sin elegir un ancho de banda. Para KDE se usa la regla de Scott y se identifica la transformación en el eje.

# %%
def density(ax,values,label,color):
    x=np.asarray(pd.Series(values).dropna(),dtype=float)
    if len(x)>=3 and np.unique(x).size>=2:
        grid=np.linspace(x.min(),x.max(),220)
        ax.plot(grid,gaussian_kde(x,bw_method='scott')(grid),label=f'{label} (n={len(x)})',color=color,lw=2)
    else: ax.text(.05,.85,'Datos insuficientes para KDE',transform=ax.transAxes)
fig,axes=plt.subplots(1,3,figsize=(15,4.3))
axes[0].hist(pme.estimacion_clp/1e6,bins=35,color=COLORS[0]); axes[0].axvline(pme.estimacion_clp.median()/1e6,c=COLORS[1],label='Mediana')
axes[0].set(title=f'PME 2024: {len(pme)} acciones',xlabel='Millones CLP por acción',ylabel='Frecuencia'); axes[0].legend()
for (sector,g),c in zip(pme.groupby('sector'),COLORS):
    density(axes[1],np.log10(g.loc[g.estimacion_clp.gt(0),'estimacion_clp']),sector,c)
axes[1].set(title='Densidad de importes positivos',xlabel='log10(CL P) por acción'.replace('CL P','CLP'),ylabel='Densidad en escala log10'); axes[1].legend(fontsize=8)
x=np.sort(pme.estimacion_clp.to_numpy()); axes[2].step(x,np.arange(1,len(x)+1)/len(x),where='post',color=COLORS[0])
axes[2].set_xscale('symlog',linthresh=1000)
axes[2].set_xticks([0,1e5,1e6,1e7,1e8],['0','100 mil','1 M','10 M','100 M'])
axes[2].yaxis.set_major_formatter(PercentFormatter(1))
axes[2].set(title='ECDF, incluyendo estimaciones cero',xlabel='CLP (escala symlog)',ylabel='Fracción acumulada')
finish(fig,'03_distribucion_pme')
table(pme.groupby('sector').agg(acciones=('fila_excel','size'),cero=('estimacion_cero','sum'),mediana=('estimacion_clp','median'))
      .assign(pct_cero=lambda d:100*d.cero/d.acciones).reset_index(),'pme_ceros_por_sector')
note('Las KDE se normalizan por sector: una curva más alta no significa más escuelas. Las acciones de un mismo colegio pueden parecerse; '
     'esta figura describe acciones, no una muestra aleatoria de colegios. La ECDF hace visible la masa de ceros que excluimos del suavizado.')

# %% [markdown]
# ### 7.2 OC: distribuciones por año, estado y moneda
#
# Se usa una OC por código. Las figuras principales monetarias incluyen solo estados Aceptada/Recepción Conforme en CLP. Otras monedas quedan fuera porque no tenemos una conversión fechada. Se informa cuántas se excluyen. Ni las órdenes aceptadas ni la recepción conforme acreditan por sí solas pago efectivo.
#
# **2022:** falta Compra Ágil del segundo semestre en la descarga municipal. **2025:** solo OC inequívocas de colegios de La Cisterna dentro del SLEP; no todas las compras del SLEP ni de sus ocho colegios. No se interpreta una caída 2024–2025 como ahorro o menor gasto.

# %%
oc_included=oc.loc[oc.monto_oc_vigente_clp.notna()].copy()
oc_included['log10_monto']=np.log10(oc_included.monto_oc_vigente_clp.where(oc_included.monto_oc_vigente_clp.gt(0)))
table(oc.groupby('anio').agg(oc_seleccionadas=('codigo_oc','size'),oc_con_rbd=('rbd_asignado','count'),
                            monto_total_clp_disponible=('monto_total_oc_clp','count'),oc_monetarias_incluidas=('monto_oc_vigente_clp','count')).reset_index(),'oc_denominadores_eda')
fig,axes=plt.subplots(2,2,figsize=(12,8))
for (year,g),ax in zip(oc_included.groupby('anio'),axes.flat):
    x=g.log10_monto.dropna()
    ax.hist(x,bins=min(20,max(3,int(np.sqrt(len(x))))),density=True,alpha=.45,color=COLORS[0])
    density(ax,x,str(year),COLORS[0])
    ax.set(title=f'{year}: {len(g)} OC elegibles; {len(x)} positivas',xlabel='log10(CLP) por OC',ylabel='Densidad')
    ax.set_xlim(3,9)
    ax.legend(fontsize=8)
finish(fig,'04_distribucion_oc_por_anio')
note('Cada panel normaliza su distribución. El tamaño reducido de 2025 hace que su densidad sea inestable. '
     'Se muestran los puntos/extremos y las medidas robustas en la sección de revisión; no se suaviza la falta de cobertura.')

# %% [markdown]
# ### 7.3 SIMCE: histogramas y densidades sin mezclar aplicaciones
#
# Se presentan Lectura y Matemática por cada combinación año–grado observada. Historia de 8° básico 2025 se muestra aparte. Cada leyenda incluye el número de puntajes válidos. No se imputan ausentes. Las densidades aproximan distribuciones de **puntajes medios por establecimiento**, no puntajes individuales de estudiantes.

# %%
applications=sorted(simce[['anio','grado']].drop_duplicates().itertuples(index=False,name=None))
fig,axes=plt.subplots(4,2,figsize=(13,14),sharex=True)
for (year,grade),ax in zip(applications,axes.flat):
    g=simce.loc[simce.anio.eq(year)&simce.grado.eq(grade)]
    for area,color in [('lect',COLORS[0]),('mate',COLORS[1])]:
        x=g.loc[g.area.eq(area),'puntaje'].dropna()
        ax.hist(x,bins=np.arange(170,371,20),density=True,alpha=.22,color=color)
        density(ax,x,{'lect':'Lectura','mate':'Matemática'}[area],color)
    ax.set(title=f'{year} · {grade}',xlabel='Puntaje SIMCE del establecimiento',ylabel='Densidad'); ax.legend(fontsize=8)
finish(fig,'05_simce_distribuciones_aplicacion')
hist=simce.loc[simce.area.eq('hist')]
fig,ax=plt.subplots(figsize=(8,3.8)); ax.hist(hist.puntaje.dropna(),bins=12,density=True,alpha=.35,color=COLORS[2])
density(ax,hist.puntaje,'Historia 8b 2025',COLORS[2]); ax.legend()
ax.set(xlabel='Puntaje SIMCE Historia',ylabel='Densidad',title='Historia se conserva como área independiente')
finish(fig,'06_simce_historia')
table(simce.groupby(['anio','grado','version_base']).agg(filas_area=('rbd','size'),puntajes_validos=('puntaje','count')).reset_index(),'simce_versiones')
note('Los resultados se separan por aplicación para evitar atribuir diferencias a grados o cohortes distintos. '
     'Versiones preliminares y finales se identifican en la tabla. Los valores faltantes y las escuelas sin fila se examinan a continuación.')

# %% [markdown]
# ### 7.4 IDPS: indicadores y versiones
#
# AM = autoestima académica y motivación; CC = clima de convivencia; PF = participación y formación ciudadana; HV = hábitos de vida saludable. Se muestran distribuciones 2024 por grado, cada indicador en su propio panel. El rango 0–100 describe la escala, no una meta de aprobación. El resto de años queda en tablas y una figura por aplicación del indicador CC para observar variación sin colapsar grados.

# %%
fig,axes=plt.subplots(2,2,figsize=(12,8),sharex=True)
for indicator,ax in zip(['AM','CC','PF','HV'],axes.flat):
    for (grade,g),color in zip(idps.loc[idps.anio.eq(2024)&idps.indicador_std.eq(indicator)].groupby('grado'),COLORS):
        x=g.puntaje.dropna()
        ax.hist(x,bins=np.arange(40,101,5),density=True,alpha=.15,color=color)
        density(ax,x,grade,color)
    ax.set(title=indicator,xlabel='Puntaje medio IDPS del establecimiento',ylabel='Densidad',xlim=(40,100)); ax.legend(fontsize=8)
finish(fig,'07_idps_2024_indicadores')
fig,axes=plt.subplots(4,2,figsize=(12,13),sharex=True)
for (year,grade),ax in zip(applications,axes.flat):
    x=idps.loc[idps.anio.eq(year)&idps.grado.eq(grade)&idps.indicador_std.eq('CC'),'puntaje'].dropna()
    ax.hist(x,bins=np.arange(40,101,5),density=True,alpha=.4,color=COLORS[2]); density(ax,x,'CC',COLORS[2])
    ax.set(title=f'Clima de convivencia {year} · {grade}',xlabel='Puntaje IDPS',ylabel='Densidad',xlim=(40,100)); ax.legend(fontsize=8)
finish(fig,'08_idps_cc_aplicaciones')
table(idps.groupby(['anio','grado','version_base']).agg(filas=('rbd','size'),puntajes_validos=('puntaje','count')).reset_index(),'idps_versiones')

# %% [markdown]
# ## 8. Missingness: dónde falta información y qué significa
#
# Un mapa de ausencias evita ocultar quién queda fuera del análisis. Hay que distinguir:
#
# - **Valor ausente dentro de una fila publicada:** la fila existe, el puntaje no.
# - **Sin fila en el extracto:** puede depender de oferta, aplicación o cobertura; sin el padrón de elegibilidad no se etiqueta automáticamente como incumplimiento.
# - **Fuera del período/fuente recopilada:** PME distinto de 2024, Agencia 2022, OC de particulares.
# - **Cero válido:** una estimación cero, o cero filas vinculadas dentro del extracto.
#
# Los porcentajes entre filas publicadas usan esa base como denominador. Los del panel usan RBD–año del censo: contestan otra pregunta. Las escuelas especiales permanecen en el censo y requieren indicadores pertinentes; no se les inventa SIMCE cero.

# %%
simce_missing=simce.assign(falta=lambda d:d.puntaje.isna()).groupby(['anio','grado','area']).agg(
    filas=('rbd','size'),faltantes=('falta','sum')).reset_index()
simce_missing['pct_faltante']=100*simce_missing.faltantes/simce_missing.filas
idps_missing=idps.assign(falta=lambda d:d.puntaje.isna()).groupby(['anio','grado','indicador_std']).agg(
    filas=('rbd','size'),faltantes=('falta','sum')).reset_index()
idps_missing['pct_faltante']=100*idps_missing.faltantes/idps_missing.filas
table(simce_missing,'missing_simce_por_aplicacion')
table(idps_missing.loc[idps_missing.faltantes.gt(0)],'missing_idps_con_ausencias')
idps_missing.to_csv(OUT/'missing_idps_todas_aplicaciones.csv',index=False)

def heatmap(ax,frame,title,fmt='.1f',cmap='Blues',vmin=0,vmax=None):
    arr=frame.to_numpy(dtype=float)
    im=ax.imshow(np.ma.masked_invalid(arr),aspect='auto',cmap=cmap,vmin=vmin,vmax=vmax)
    ax.set_xticks(range(len(frame.columns)),frame.columns,rotation=45,ha='right')
    ax.set_yticks(range(len(frame.index)),frame.index)
    ax.set_title(title)
    for i in range(arr.shape[0]):
        for j in range(arr.shape[1]):
            if np.isfinite(arr[i,j]):
                ax.text(j,i,format(arr[i,j],fmt),ha='center',va='center',fontsize=9,
                        color='white' if vmax and arr[i,j]>vmax*.65 and vmin==0 else '#203248')
    return im
sm=simce_missing.assign(aplicacion=lambda d:d.anio.astype(str)+' '+d.grado).pivot(index='aplicacion',columns='area',values='pct_faltante')
im=idps_missing.assign(aplicacion=lambda d:d.anio.astype(str)+' '+d.grado).pivot(index='aplicacion',columns='indicador_std',values='pct_faltante')
fig,axes=plt.subplots(1,2,figsize=(12,6))
heatmap(axes[0],sm,'% sin puntaje entre filas SIMCE publicadas',vmax=100)
heatmap(axes[1],im,'% sin puntaje entre filas IDPS publicadas',vmax=100)
finish(fig,'09_missingness_agencia_por_aplicacion')
note('Una celda vacía de Historia fuera de 8b 2025 significa que esa área no existe en la aplicación disponible, '
     'no que el 100% de alumnos tenga un puntaje faltante. Los ceros de la tabla indican 0% de ausencia entre filas, no puntaje cero.')

# %% [markdown]
# ### 8.1 Matriz por colegio: ausencia de fila y ausencia de valor
#
# Usamos 4° básico 2024 como corte concreto, sin asumir que los 59 colegios debieron rendirlo. Gris significa sin fila, amarillo fila sin valor y azul valor publicado. No se suman porcentajes entre indicadores.

# %%
census24=directory.loc[directory.anio.eq(2024)].sort_values(['oferta','sector','rbd']).set_index('rbd')
status=pd.DataFrame(index=census24.index)
for label,df,key in [('SIMCE Lectura',simce.loc[simce.anio.eq(2024)&simce.grado.eq('4b')&simce.area.eq('lect')],'puntaje'),
                      ('SIMCE Matemática',simce.loc[simce.anio.eq(2024)&simce.grado.eq('4b')&simce.area.eq('mate')],'puntaje')]:
    mapping=df.set_index('rbd')[key]
    status[label]=[2 if r in mapping.index and pd.notna(mapping.loc[r]) else 1 if r in mapping.index else 0 for r in status.index]
for indicator in ['AM','CC','PF','HV']:
    mapping=idps.loc[idps.anio.eq(2024)&idps.grado.eq('4b')&idps.indicador_std.eq(indicator)].set_index('rbd').puntaje
    status['IDPS '+indicator]=[2 if r in mapping.index and pd.notna(mapping.loc[r]) else 1 if r in mapping.index else 0 for r in status.index]
fig,ax=plt.subplots(figsize=(10,14))
ax.imshow(status,aspect='auto',cmap=ListedColormap(['#dce1e7','#e3ac4c','#176b87']),vmin=0,vmax=2)
ax.set_xticks(range(len(status.columns)),status.columns,rotation=30,ha='right')
ax.set_yticks(range(len(status)),[f'{r} · {census24.loc[r,"oferta"]}' for r in status.index],fontsize=8)
ax.set_title('4° básico 2024: gris = sin fila; amarillo = sin valor; azul = valor')
finish(fig,'10_missingness_agencia_por_rbd')
table(status.apply(pd.Series.value_counts).fillna(0).rename(index={0:'Sin fila',1:'Fila sin puntaje',2:'Con puntaje'}).reset_index(names='estado'),'missing_4b_2024_estados')

# %% [markdown]
# ### 8.2 Mapa del panel y mecanismos de ausencia
#
# Se muestran columnas sustantivas y se segmentan las tasas por año y oferta. Esta visualización revela bloques estructurales: PME solo 2024, Agencia desde 2023 y OC solo sector público. Por eso una imputación masiva a cero o mediana alteraría el significado de la investigación. No diagnosticamos MCAR/MAR/MNAR solo por mirar el mapa.

# %%
missing_cols=['matricula_total','pas_filas_archivo_anual','pme_estimacion_clp_2024','pme_pct_senal_revision',
              'simce_4b_lect','simce_4b_mate','simce_2m_lect','idps_4b_AM','idps_4b_CC',
              'oc_rbd_vinculadas','oc_monto_vigente_clp_rbd_vinculado']
missing_labels=['Matrícula','Filas PAS','Estimación PME','KPI PME','SIMCE 4b Lect.','SIMCE 4b Mat.',
                'SIMCE 2m Lect.','IDPS 4b AM','IDPS 4b CC','OC vinculadas','Monto OC CLP']
p_sort=panel.sort_values(['anio','sector','oferta','rbd']).reset_index(drop=True)
fig,axes=plt.subplots(1,2,figsize=(14,7),gridspec_kw={'width_ratios':[1.6,1]})
axes[0].imshow(p_sort[missing_cols].isna(),aspect='auto',cmap=ListedColormap(['#176b87','#eeeeee']),vmin=0,vmax=1)
axes[0].set_xticks(range(len(missing_cols)),missing_labels,rotation=60,ha='right')
axes[0].set(title='Panel: azul = observado; gris = ausente',ylabel='Filas RBD–año ordenadas')
for year,g in p_sort.groupby('anio'):
    axes[0].axhline(g.index.min()-.5,color='#d28e32',lw=1)
    axes[0].text(-.8,np.mean(g.index),str(year),ha='right',fontsize=9)
rates=panel[missing_cols].isna().mean()*100
axes[1].barh(missing_labels,rates,color=COLORS[0]); axes[1].set(xlabel='% de 238 RBD–año sin valor',xlim=(0,100),title='Ausencia total, con causas diferentes')
finish(fig,'11_missingness_panel')
table(panel.assign(falta_pme=panel.pme_pct_senal_revision.isna(),falta_simce4b=panel.simce_4b_lect.isna(),
                   falta_idps4b=panel.idps_4b_CC.isna(),falta_monto_oc=panel.oc_monto_vigente_clp_rbd_vinculado.isna())
      .groupby(['anio','sector','oferta']).agg(rbd=('rbd','size'),sin_kpi_pme=('falta_pme','sum'),
                                             sin_simce4b=('falta_simce4b','sum'),sin_idps4b=('falta_idps4b','sum'),sin_monto_oc=('falta_monto_oc','sum')).reset_index(),
      'missing_panel_por_segmento')
coverage_reasons=pd.DataFrame([
    ['PME 2022/2023/2025',int(panel.anio.ne(2024).sum()),'Año PME no recopilado'],
    ['PME 2024 sin fila',int((panel.anio.eq(2024)&panel.pme_pct_senal_revision.isna()).sum()),'Verificar aplicabilidad/cobertura; no KPI cero'],
    ['SIMCE/IDPS 2022',int(panel.anio.eq(2022).sum()),'Año Agencia no recopilado'],
    ['OC particulares',int(panel.sector.eq('Particular subvencionado').sum()),'Fuente de compras no recopilada para este sector'],
    ['OC públicos sin vínculo',int((panel.sector.eq('Público')&panel.oc_rbd_vinculadas.eq(0)).sum()),'Cero vínculos confirmados, no cero compras']
],columns=['situacion','rbd_anio','interpretacion'])
table(coverage_reasons,'motivos_ausencia_panel')

# %% [markdown]
# ## 9. Outliers: identificar, revisar y medir sensibilidad
#
# Usamos cercas **Q1−1,5×IQR y Q3+1,5×IQR** como criterio descriptivo. En importes también se revisa la escala `log10(1+CLP)`: detecta qué tan dependiente es la marca de la escala. En SIMCE/IDPS las cercas se calculan dentro de año–grado–área/indicador. Los grupos pequeños no permiten una referencia estable.
#
# Una marca no prueba error, fraude ni ilegalidad. Se conservan todos los registros en el cálculo principal, se muestran localizadores reales y se propone qué respaldo revisar. El contraste «sin extremos» es una sensibilidad, no una limpieza validada.

# %%
def iqr_flags(s):
    x=pd.to_numeric(s,errors='raise').astype(float)
    q1,q3=x.quantile([.25,.75]); iqr=q3-q1
    low,high=q1-1.5*iqr,q3+1.5*iqr
    flags=x.lt(low)|x.gt(high)
    return flags,low,high
pme['extremo_iqr'],pme_low,pme_high=iqr_flags(pme.estimacion_clp)
pme['extremo_iqr_log'],_,_=iqr_flags(pme.log10_estimacion_mas1)
oc_included['extremo_iqr']=False
outlier_summary=[]
for year,g in oc_included.groupby('anio'):
    flags,low,high=iqr_flags(g.monto_oc_vigente_clp)
    oc_included.loc[g.index,'extremo_iqr']=flags
    outlier_summary.append({'fuente':'OC','grupo':str(year),'n':len(g),'extremos':int(flags.sum()),'limite_inferior':low,'limite_superior':high})
for name,df,groups in [('SIMCE',simce,['anio','grado','area']),('IDPS',idps,['anio','grado','indicador_std'])]:
    df['extremo_iqr']=False
    for key,g in df.groupby(groups):
        flags,low,high=iqr_flags(g.puntaje)
        df.loc[g.index,'extremo_iqr']=flags
        outlier_summary.append({'fuente':name,'grupo':' '.join(map(str,key)),'n':int(g.puntaje.notna().sum()),
                                'extremos':int(flags.sum()),'limite_inferior':low,'limite_superior':high})
table(pd.DataFrame([{'fuente':'PME','grupo':'CLP','n':len(pme),'extremos':int(pme.extremo_iqr.sum()),'limite_inferior':pme_low,'limite_superior':pme_high},
                    {'fuente':'PME','grupo':'log10(1+CLP)','n':len(pme),'extremos':int(pme.extremo_iqr_log.sum())}]+outlier_summary),'outliers_resumen')
table(pme.sort_values('estimacion_clp',ascending=False)[['fila_excel','rbd','nombre','dimension','NOM_ACTIVIDAD','estimacion_clp','estado_declarado','extremo_iqr']].head(12),'pme_extremos_revisar')
table(oc_included.sort_values('monto_oc_vigente_clp',ascending=False)[['anio','codigo_oc','nombre_oc','estado_oc','rbd_asignado','monto_oc_vigente_clp','extremo_iqr']].head(10),'oc_extremos_revisar')
agency_extremes=pd.concat([simce.loc[simce.extremo_iqr,['anio','rbd','grado','area','puntaje','version_base']].assign(fuente='SIMCE').rename(columns={'area':'medida'}),
                          idps.loc[idps.extremo_iqr,['anio','rbd','grado','indicador_std','puntaje','version_base']].assign(fuente='IDPS').rename(columns={'indicador_std':'medida'})],ignore_index=True)
table(agency_extremes,'agencia_extremos_revisar')

# %%
fig,axes=plt.subplots(1,2,figsize=(14,5))
dim_order=dimensions.dimension.tolist()
axes[0].boxplot([pme.loc[pme.dimension.eq(d),'estimacion_clp']/1e6 for d in dim_order],tick_labels=[d.replace(' ','\n',1) for d in dim_order],showfliers=True)
axes[0].set(title='PME por dimensión: todos los valores',ylabel='Estimación por acción (millones CLP)')
axes[1].boxplot([g.log10_monto.dropna() for _,g in oc_included.groupby('anio')],tick_labels=sorted(oc_included.anio.unique()),showfliers=True)
axes[1].set(title='OC elegibles por año: escala logarítmica',ylabel='log10(CLP) por OC')
finish(fig,'12_boxplots_montos')

# %% [markdown]
# ### 9.1 ¿Cuánto cambian las conclusiones si apartamos extremos?
#
# Comparamos datos completos con exclusión hipotética IQR y winsorización al P99 solo del importe. Winsorizar cambia valores; **no se publica como base limpia**. También contrastamos el KPI al excluir acciones extremas, porque el sesgo de retirar filas podría cambiar un porcentaje aunque su fórmula no use el monto continuo.

# %%
sensitivity=[]
for label,frame,amount in [('Todos los registros',pme,pme.estimacion_clp),
                           ('Sin extremos IQR (sensibilidad)',pme.loc[~pme.extremo_iqr],pme.loc[~pme.extremo_iqr,'estimacion_clp']),
                           ('Importes limitados a P99 (sensibilidad)',pme,pme.estimacion_clp.clip(upper=pme.estimacion_clp.quantile(.99)))]:
    sensitivity.append({'escenario':label,'acciones':len(frame),'media_clp':amount.mean(),'mediana_clp':amount.median(),
                        'suma_clp':amount.sum(),'pct_senal_revision':100*frame.senal_revision.mean()})
sensitivity=pd.DataFrame(sensitivity)
table(sensitivity,'sensibilidad_extremos_pme')
topshare=pme.nlargest(max(1,int(np.ceil(len(pme)*.1))),'estimacion_clp').estimacion_clp.sum()/pme.estimacion_clp.sum()
note(f'**Concentración:** el 10% superior de acciones (redondeado hacia arriba) reúne {100*topshare:.1f}% de la estimación. '
     f'La cerca IQR superior es ${esnum(pme_high)}. Hay {pme.extremo_iqr.sum()} acciones sobre/fuera de las cercas. '
     'Se solicita para ellas presupuesto desglosado, beneficiarios y explicación de escala. Los importes permanecen en la base principal.')

# %% [markdown]
# ## 10. Correlaciones del panel: diagnóstico y descripción
#
# **Corte principal: 2024, una fila por RBD.** Así evitamos repetir el mismo PME en otros años y tratar observaciones del mismo colegio como independientes. La matriz usa matrícula, PAS publicados, número de acciones, logaritmo de estimación, KPI y resultados específicos de 4° básico. No incluye RBD ni códigos de dependencia como variables numéricas.
#
# **Pearson** mide asociación lineal y puede cambiar por extremos. **Spearman** usa rangos y recoge asociaciones monótonas. Se elimina ausencia **por pareja**, por eso se muestra una matriz de `n` para cada coeficiente. Se ocultan coeficientes con menos de 10 pares o varianza cero. Es una regla de presentación prudente, no un umbral de validez estadística. La muestra completa cambia entre parejas.
#
# No presentamos p-valores ni conclusiones causales: las escuelas observadas no son una muestra aleatoria; oferta, matrícula, sector y disponibilidad condicionan las asociaciones. El número de acciones y el total estimado comparten una relación de construcción. OC se analiza aparte: en 2024 solo hay ocho públicos y una atribución parcial.

# %%
p24=panel.loc[panel.anio.eq(2024)].copy()
corr_vars=['matricula_total','pas_filas_archivo_anual','pme_acciones','pme_log10_total_mas1','pme_pct_senal_revision',
           'simce_4b_lect','simce_4b_mate','idps_4b_AM','idps_4b_CC']
corr_names=['Matrícula','PAS publicados','Acciones PME','log10(PME+1)','KPI revisión %','SIMCE 4b Lect.','SIMCE 4b Mat.','IDPS 4b AM','IDPS 4b CC']
cdata=p24[corr_vars].astype(float)
counts=cdata.notna().astype(int).T@cdata.notna().astype(int)
pearson=cdata.corr(method='pearson',min_periods=10)
spearman=cdata.corr(method='spearman',min_periods=10)
for matrix,name in [(pearson,'correlacion_pearson_2024'),(spearman,'correlacion_spearman_2024'),(counts,'correlacion_n_pares_2024')]:
    matrix.to_csv(OUT/f'{name}.csv')
fig,axes=plt.subplots(1,2,figsize=(17,7))
for ax,matrix,title in [(axes[0],pearson,'Pearson · 2024'),(axes[1],spearman,'Spearman · 2024')]:
    data=matrix.copy(); data.index=corr_names; data.columns=corr_names
    heatmap(ax,data,title,fmt='.2f',cmap='coolwarm',vmin=-1,vmax=1)
finish(fig,'13_correlaciones_panel_2024')
fig,ax=plt.subplots(figsize=(10,8)); nc=counts.copy(); nc.index=corr_names; nc.columns=corr_names
heatmap(ax,nc,'Número de colegios observados por pareja',fmt='.0f',vmax=59)
finish(fig,'14_correlaciones_numero_pares')
note('Las asociaciones describen el subconjunto con cada par disponible. Un coeficiente entre KPI y SIMCE '
     'no verifica implementación ni eficiencia del dinero. No hay una secuencia temporal de exposición comparable ni un grupo de control.')

# %% [markdown]
# ### 10.1 Dispersión, segmentos y sensibilidad de las asociaciones
#
# Revisamos tres pares interpretables, mostrando sector y tamaño de las muestras. La asociación matrícula–estimación puede reflejar escala del colegio; la de estimación–SIMCE puede reflejar composición del grupo. La tabla por sector y oferta permite detectar si el agregado oculta diferencias; no afirmamos paradoja de Simpson sin una inversión observada.

# %%
pairs=[('matricula_total','pme_log10_total_mas1','Matrícula total','log10(estimación PME + 1)'),
       ('pme_log10_total_mas1','simce_4b_lect','log10(estimación PME + 1)','SIMCE Lectura 4b'),
       ('pme_pct_senal_revision','idps_4b_CC','KPI de revisión (%)','IDPS clima convivencia 4b')]
fig,axes=plt.subplots(1,3,figsize=(16,4.6))
for ax,(x,y,xlabel,ylabel) in zip(axes,pairs):
    for (sector,g),color in zip(p24.groupby('sector'),COLORS):
        obs=g[[x,y]].dropna(); ax.scatter(obs[x],obs[y],label=f'{sector} (n={len(obs)})',color=color,alpha=.75,s=35)
    ax.set(xlabel=xlabel,ylabel=ylabel,title='Corte 2024'); ax.legend(fontsize=7)
finish(fig,'15_dispersion_segmentada')
assoc=[]
segments=[('Todos',p24)]+[(s,g) for s,g in p24.groupby('sector')]+[(o,g) for o,g in p24.groupby('oferta')]
for segment,g in segments:
    for x,y,_,_ in pairs:
        d=g[[x,y]].dropna(); enough=len(d)>=10 and d[x].nunique()>1 and d[y].nunique()>1
        assoc.append({'segmento':segment,'x':x,'y':y,'n':len(d),
                      'Pearson':d[x].corr(d[y]) if enough else np.nan,
                      'Spearman':d[x].corr(d[y],method='spearman') if enough else np.nan,
                      'Kendall':d[x].corr(d[y],method='kendall') if enough else np.nan,
                      'nota':'Descriptiva' if enough else 'No reportado: n<10 o sin variación'})
table(pd.DataFrame(assoc),'asociaciones_segmentadas')
for x,y,_,_ in pairs:
    d=p24[[x,y]].dropna()
    if len(d)>=10 and d[x].nunique()>1 and d[y].nunique()>1:
        note(f'**{x} / {y}:** n={len(d)}, Pearson={d[x].corr(d[y]):.2f}, '
             f'Spearman={d[x].corr(d[y],method="spearman"):.2f}. '
             'La comparación se limita a esos RBD; la tabla por sector/oferta muestra su composición.')
school_amount=p24.pme_estimacion_clp_2024.dropna()
school_flags,_,_=iqr_flags(school_amount)
top3=school_amount.nlargest(3).index
scenarios=[('Todos los colegios con PME',p24),
           (f'Sin colegios extremos por regla IQR ({int(school_flags.sum())} marcados)',p24.drop(index=school_flags.index[school_flags])),
           ('Sin los 3 colegios de mayor importe',p24.drop(index=top3))]
assoc_sensitivity=[]
for label,d in scenarios:
    obs=d[['matricula_total','pme_estimacion_clp_2024']].dropna()
    assoc_sensitivity.append({'escenario':label,'n':len(obs),
                              'Pearson_CLP':obs.matricula_total.corr(obs.pme_estimacion_clp_2024),
                              'Pearson_log10':obs.matricula_total.corr(np.log10(obs.pme_estimacion_clp_2024+1)),
                              'Spearman':obs.matricula_total.corr(obs.pme_estimacion_clp_2024,method='spearman')})
table(pd.DataFrame(assoc_sensitivity),'sensibilidad_correlacion_extremos')
note(f'**Sensibilidad matrícula–estimación PME:** la regla IQR aplicada a los totales por colegio marca **{int(school_flags.sum())}** colegios, '
     'por lo que ese escenario coincide con el completo. Como control más exigente se retiran los tres colegios de mayor importe. '
     'Pearson se reporta en CLP y en log10 porque la escala cambia su valor; Spearman, basado en rangos, no depende de esa escala.')
public24=p24.loc[p24.sector.eq('Público'),['rbd','matricula_total','pme_estimacion_clp_2024','oc_rbd_vinculadas','oc_monto_vigente_clp_rbd_vinculado']]
table(public24,'oc_pme_ocho_publicos_2024')
note('Las ocho escuelas públicas no se mezclan con particulares con OC ausente para calcular una correlación monetaria. '
     'Los montos OC y PME tienen universos distintos: no se restan ni se suman para construir supuestas brechas de gasto.')

# %% [markdown]
# ## 11. Respuesta por dimensión y establecimiento
#
# ### 11.1 Acciones, implementación completa, ceros y mediana por dimensión
#
# Las cuatro medidas usan las mismas acciones. Las tasas se ponderan por acción; una media simple de tasas de colegios contestaría otra pregunta. La mediana incluye estimaciones cero. El orden se mantiene en los cuatro gráficos para facilitar comparación.

# %%
dim_order=['Convivencia Escolar','Gestión Pedagógica','Gestión de Recursos','Liderazgo']
dims=dimensions.set_index('dimension').loc[dim_order]
fig,axes=plt.subplots(2,2,figsize=(13,8))
for ax,col,label,color in [(axes[0,0],'acciones','Acciones observadas',COLORS[0]),
                            (axes[0,1],'pct_completas','Declaradas completas (%)',COLORS[2]),
                            (axes[1,0],'pct_estimacion_cero','Estimación $0 (%)',COLORS[1]),
                            (axes[1,1],'estimacion_mediana_clp','Mediana por acción (millones CLP)',COLORS[3])]:
    values=dims[col]/1e6 if col=='estimacion_mediana_clp' else dims[col]
    ax.barh(dim_order,values,color=color); ax.invert_yaxis(); ax.set_title(label)
    ax.set_xlim(0,values.max()*1.25)
    for i,v in enumerate(values): ax.text(v+values.max()*.02,i,f'{v:.1f}' if col!='acciones' else str(int(v)),va='center')
finish(fig,'16_pme_dimensiones')
table(dimensions[['dimension','acciones','completas','pct_completas','estimacion_cero','pct_estimacion_cero','estimacion_mediana_clp','senal_revision','pct_senal_revision']],'dimension_resultados')
highest=dimensions.sort_values('pct_senal_revision',ascending=False).iloc[0]
note(f'**Mayor proporción de señales:** {highest.dimension}, {int(highest.senal_revision)}/{int(highest.acciones)} '
     f'({highest.pct_senal_revision:.1f}%). Este orden describe el registro. Gestión de Recursos puede tener una mediana mayor '
     'por la naturaleza de sus acciones; una dimensión más costosa no demuestra ineficiencia.')

# %% [markdown]
# ### 11.2 KPI por colegio y por colegio–dimensión
#
# La figura incluye los 46 colegios con acciones y muestra **señales/acciones** al lado de cada barra. El mapa por dimensión permite distinguir una señal generalizada de una concentrada en una dimensión. Una celda ausente significa ninguna acción de esa dimensión en el extracto, no 0% de señal.

# %%
ranked=schools.sort_values(['pct_senal_revision','senal_revision','rbd'],ascending=[False,False,True]).reset_index(drop=True)
fig,ax=plt.subplots(figsize=(12,15))
ax.barh(range(len(ranked)),ranked.pct_senal_revision,color=[COLORS[1] if s=='Público' else COLORS[0] for s in ranked.sector])
ax.set_yticks(range(len(ranked)),[f'{r.rbd} · {r.nombre[:38]}' for r in ranked.itertuples()],fontsize=8)
ax.invert_yaxis(); ax.axvline(50,color='#933d48',ls='--',lw=1.5,label='Umbral operativo propuesto 50%')
ax.set(xlim=(0,120),xlabel='Acciones con señal / acciones observadas (%)',title='PME 2024 por colegio: naranja público, azul particular subvencionado')
for i,r in ranked.iterrows(): ax.text(r.pct_senal_revision+1,i,f'{int(r.senal_revision)}/{int(r.acciones)}',va='center',fontsize=8)
ax.legend(loc='lower right',fontsize=8)
finish(fig,'17_kpi_todos_colegios')
sd_rate=school_dimensions.pivot(index='rbd',columns='dimension',values='pct_senal_revision').reindex(index=ranked.rbd,columns=dim_order)
sd_num=school_dimensions.pivot(index='rbd',columns='dimension',values='senal_revision').reindex_like(sd_rate)
sd_den=school_dimensions.pivot(index='rbd',columns='dimension',values='acciones').reindex_like(sd_rate)
fig,ax=plt.subplots(figsize=(10,15))
ax.imshow(np.ma.masked_invalid(sd_rate.to_numpy(dtype=float,na_value=np.nan)),aspect='auto',cmap='YlOrRd',vmin=0,vmax=100)
ax.set_xticks(range(4),[x.replace(' ','\n',1) for x in dim_order]); ax.set_yticks(range(len(ranked)),ranked.rbd,fontsize=8)
ax.set(title='Colegio–dimensión: color = % de señal; etiqueta = señales/acciones',ylabel='RBD, mismo orden que gráfico anterior')
for i in range(len(sd_rate)):
    for j in range(4):
        if pd.notna(sd_rate.iloc[i,j]):
            ax.text(j,i,f'{int(sd_num.iloc[i,j])}/{int(sd_den.iloc[i,j])}',ha='center',va='center',fontsize=8,
                    color='white' if sd_rate.iloc[i,j]>70 else '#203248')
finish(fig,'18_kpi_colegio_dimension')

# %% [markdown]
# ### 11.3 Composición de grupos y robustez del umbral
#
# Comparamos sectores y ofertas mostrando tanto el KPI ponderado por acciones como la mediana del KPI entre colegios. Ninguna diferencia se atribuye causalmente al tipo de sostenedor. El grupo PME contiene 46 de 59 RBD del censo 2024 y no garantiza representación del resto.
#
# La señal cero puede incluir acciones realizadas con recursos existentes o estimaciones no reflejadas en ese campo. La señal no completa también incluye avances de 75–99%. Se presenta una variante estricta «estimación cero o implementación no efectuada 0%» para comprobar cuánto depende la prioridad de la definición adoptada.

# %%
sector_kpi=aggregate_pme(['sector','oferta'])
school_median=schools.groupby(['sector','oferta']).agg(colegios=('rbd','size'),mediana_kpi_colegio=('pct_senal_revision','median')).reset_index()
table(sector_kpi.merge(school_median,on=['sector','oferta'],validate='one_to_one'),'kpi_por_sector_oferta')
threshold_rows=[]
for threshold in [25,50,75]:
    selected=schools.loc[schools.pct_senal_revision.ge(threshold)]
    threshold_rows.append({'umbral_pct':threshold,'colegios_priorizados':len(selected),
                           'de_colegios_observados':len(schools),'acciones_senaladas_en_esos_colegios':int(selected.senal_revision.sum()),
                           'publicos':int(selected.sector.eq('Público').sum()),
                           'particulares':int(selected.sector.eq('Particular subvencionado').sum())})
thresholds=pd.DataFrame(threshold_rows)
table(thresholds,'sensibilidad_umbrales')
pme['senal_estricta']=pme.estimacion_cero|pme.estado_declarado.eq('Implementación no efectuada: 0%')
strict=pme.groupby('rbd').senal_estricta.mean().mul(100)
ranked['pct_senal_estricta']=ranked.rbd.map(strict)
table(pd.DataFrame([
    {'definicion':'Cero o no declarada completa','acciones_senaladas':int(pme.senal_revision.sum()),'pct':100*pme.senal_revision.mean(),'colegios_al_50':int(schools.pct_senal_revision.ge(50).sum())},
    {'definicion':'Cero o declarada no efectuada (0%)','acciones_senaladas':int(pme.senal_estricta.sum()),'pct':100*pme.senal_estricta.mean(),'colegios_al_50':int(strict.ge(50).sum())}
]),'sensibilidad_definicion_kpi')
priority=ranked.loc[ranked.pct_senal_revision.ge(50)].sort_values(['senal_revision','pct_senal_revision','rbd'],ascending=[False,False,True]).copy()
priority['accion_propuesta']='Solicitar detalle de estimación, avance y respaldo de acciones señaladas'
table(priority[['rbd','nombre','sector','acciones','senal_revision','pct_senal_revision','estimacion_cero','no_completas','pct_senal_estricta','accion_propuesta']], 'prioridad_documental_umbral_50')
without_pme=census24.loc[~census24.index.isin(schools.rbd)].reset_index()[['rbd','nombre','sector','oferta']]
without_pme['accion_propuesta']='Verificar aplicabilidad PME y extracción; KPI no calculable'
table(without_pme,'colegios_sin_pme_verificar_cobertura')

# %% [markdown]
# ### 11.3b ¿Qué hay detrás del KPI? Estructura, composición y confundentes
#
# Antes de usar el KPI para priorizar conviene saber **de qué está hecho**. Se examinan cuatro cosas: (1) la forma de su distribución entre colegios, (2) cuánto aporta cada componente ($0 vs. no completa), (3) qué estados declarados forman «no completa» y (4) si el tipo de establecimiento o su tamaño explican parte de la señal. Son descripciones del extracto, no causas.

# %%
pme['componente']=np.select([pme.estimacion_cero&pme.no_completa,pme.estimacion_cero&~pme.no_completa,~pme.estimacion_cero&pme.no_completa],
                            ['Ambas','Solo estimación $0','Solo no completa'],'Sin señal')
composition=pd.crosstab(pme.oferta,pme.componente,margins=True,margins_name='Total')
table(composition.reset_index(),'kpi_composicion_senal')
states=pme.loc[pme.no_completa,'estado_declarado'].value_counts().rename_axis('estado_declarado').reset_index(name='acciones')
states['pct_de_no_completas']=100*states.acciones/states.acciones.sum()
table(states,'kpi_estados_no_completas')
low=int(schools.pct_senal_revision.le(25).sum()); high=int(schools.pct_senal_revision.ge(75).sum()); mid=len(schools)-low-high
zero_by_offer=pme.groupby('oferta').estimacion_cero.mean().mul(100)
regular_actions=pme.loc[pme.oferta.ne('especial')]; regular_schools=schools.loc[schools.oferta.ne('especial')]
rho_size=p24[['matricula_total','pme_pct_senal_revision']].dropna().corr(method='spearman').iloc[0,1]
advanced=int(pme.estado_declarado.eq('Implementación avanzada: 75% a 99%').sum())
structure=pd.DataFrame([
    ('colegios con KPI ≤25 %',low),('colegios con KPI entre 25 % y 75 %',mid),('colegios con KPI ≥75 %',high),
    ('acciones no completas en estado «avanzada 75–99 %»',advanced),('acciones no completas en total',int(pme.no_completa.sum())),
    ('% de acciones con $0 en escuelas especiales',float(zero_by_offer['especial'])),
    ('% de acciones con $0 en escuelas básica/media',float(zero_by_offer['basica_media'])),
    ('KPI global sin escuelas especiales (%)',float(100*regular_actions.senal_revision.mean())),
    ('colegios básica/media con KPI ≥50 %',int(regular_schools.pct_senal_revision.ge(50).sum())),
    ('colegios básica/media observados',len(regular_schools)),
    ('Spearman KPI del colegio vs. matrícula',float(rho_size))],columns=['medida','valor'])
table(structure,'kpi_estructura_resumen')
fig,axes=plt.subplots(1,2,figsize=(14,4.8))
bins=np.arange(0,101,10)
axes[0].hist([schools.loc[schools.oferta.ne('especial'),'pct_senal_revision'],schools.loc[schools.oferta.eq('especial'),'pct_senal_revision']],
             bins=bins,stacked=True,color=[COLORS[0],COLORS[3]],label=['Básica/media','Especial'],edgecolor='white')
for t in [25,50,75]: axes[0].axvline(t,color='#933d48',ls='--',lw=1)
axes[0].set(xlabel='KPI del colegio (% de acciones con señal)',ylabel='Número de colegios',title=f'KPI por colegio (n={len(schools)}): {low} ≤25 %, {mid} intermedios, {high} ≥75 %')
axes[0].legend()
share=composition.drop(index='Total',columns='Total'); share=share.div(share.sum(axis=1),axis=0).mul(100)
order=['Sin señal','Solo no completa','Solo estimación $0','Ambas']
left=np.zeros(len(share)); labels={'basica_media':'Básica/media','especial':'Especial'}
for comp,color in zip(order,['#D9E2E6',COLORS[1],COLORS[2],'#933d48']):
    axes[1].barh([labels[i] for i in share.index],share[comp],left=left,color=color,label=comp,edgecolor='white'); left=left+share[comp].to_numpy()
for i,offer in enumerate(share.index):
    axes[1].text(101,i,f'n={int(composition.loc[offer,"Total"])}',va='center',fontsize=9)
axes[1].set(xlim=(0,112),xlabel='% de las acciones del grupo',title='Composición de la señal por tipo de oferta'); axes[1].legend(ncol=4,fontsize=8,loc='upper center',bbox_to_anchor=(0.5,-0.16),frameon=False)
finish(fig,'19_kpi_estructura')
note(f'''**Lectura de la estructura del KPI** (descriptiva; no explica causas):

- **Forma bimodal:** de {len(schools)} colegios, {low} tienen KPI ≤25 % y {high} tienen ≥75 %; solo {mid} quedan en el tramo intermedio. Los colegios tienden a declarar casi todo completo o casi nada, por eso los umbrales 25/50/75 % cambian la lista de forma gradual y no abrupta.
- **«No completa» es sobre todo «casi completa»:** {advanced} de las {int(pme.no_completa.sum())} acciones no completas ({100*advanced/pme.no_completa.sum():.0f} %) están en «avanzada 75–99 %». Solo {int(pme.estado_declarado.eq('Implementación no efectuada: 0%').sum())} declaran «no efectuada 0 %». Esto explica que la definición estricta baje el KPI de {100*pme.senal_revision.mean():.1f} % a {100*pme.senal_estricta.mean():.1f} %.
- **El $0 se concentra en escuelas especiales:** {zero_by_offer['especial']:.1f} % de sus acciones tienen estimación cero, frente a {zero_by_offer['basica_media']:.1f} % en básica/media. Sin ellas, el KPI global es {100*regular_actions.senal_revision.mean():.1f} % y {int(regular_schools.pct_senal_revision.ge(50).sum())} de {len(regular_schools)} colegios superan el 50 %. Hipótesis a confirmar: que su financiamiento o su forma de planificar el PME difiera; no es evidencia de irregularidad.
- **Tamaño:** el KPI del colegio se asocia negativamente con la matrícula (Spearman {rho_size:.2f}, n={int(p24[['matricula_total','pme_pct_senal_revision']].dropna().shape[0])}): los colegios más pequeños tienden a KPI más alto. Es un posible confundente de escala; no autoriza a decir que el tamaño «causa» la señal.

**Consecuencia para la decisión:** el KPI ordena una cola de revisión documental, pero con un umbral de 50 % entran {len(priority)} de {len(schools)} colegios. Para el trabajo real conviene revisar la cola **por segmento** (especial vs. básica/media) y empezar por las acciones con ambos componentes o con estado «no efectuada», donde la señal es más específica.''')

# %% [markdown]
# ### 11.4 Qué pedir y qué aprender de los casos
#
# A cada colegio priorizado se le solicitaría la acción identificada por nombre, dimensión, año y localizador, su presupuesto desglosado, explicación del cero si corresponde, medios de verificación, avance y razones de reprogramación. Facturas/pagos/recepción se pedirían cuando exista una operación asociada que los requiera. El primer paso es confirmar que acción y operación corresponden, no inferir que una OC financió un PME por compartir RBD y año.
#
# También conviene revisar una pequeña selección de colegios bajo el umbral y casos sin señal para conocer falsas alarmas y omisiones del criterio. No se estima sensibilidad/especificidad porque faltan etiquetas independientes verificadas. Los listados son propuestas de trabajo: no se envió ninguna solicitud.

# %%
evidence_cases=pme.loc[pme.rbd.isin(priority.rbd)&pme.senal_revision,
    ['rbd','nombre','fila_excel','dimension','NOM_ACTIVIDAD','DESC_ACTIVIDAD','estimacion_clp','estado_declarado','estimacion_cero','no_completa','extremo_iqr']]
evidence_cases.to_csv(OUT/'acciones_para_respaldo_documental.csv',index=False)
note(f'Con umbral ≥50% se priorizan **{len(priority)} de {len(schools)} colegios observados**. '
     f'El listado contiene {len(evidence_cases)} acciones señaladas de esos colegios. '
     f'Hay **{len(without_pme)} RBD sin acciones PME** en el extracto: requieren aclarar cobertura antes de asignarles un KPI.')

# %% [markdown]
# ## 12. Conclusiones calculadas y límites de la respuesta
#
# La respuesta se apoya en distribuciones, denominadores y asociaciones visibles, no solo en presencia de archivos. El bloque siguiente calcula los hallazgos principales para que las cifras se actualicen al repetir la ejecución. Las interpretaciones son condicionales al extracto, al corte temporal y a la regla de prioridad.

# %%
result={
    'fecha_revision':'2026-09-28','pregunta':'Distribución de señales documentales en acciones PME 2024 por colegio y dimensión',
    'acciones':len(pme),'colegios_pme':len(schools),'colegios_censo_2024':len(census24),
    'completas':int(pme.completa.sum()),'no_completas':int(pme.no_completa.sum()),'cero':int(pme.estimacion_cero.sum()),
    'ambas':int(pme.ambas.sum()),'senal_union':int(pme.senal_revision.sum()),'pct_senal_union':float(100*pme.senal_revision.mean()),
    'media_estimacion_clp':float(pme.estimacion_clp.mean()),'mediana_estimacion_clp':float(pme.estimacion_clp.median()),
    'extremos_pme_iqr':int(pme.extremo_iqr.sum()),'sin_puntaje_simce':int(simce.puntaje.isna().sum()),
    'medidas_simce':len(simce),'sin_puntaje_idps':int(idps.puntaje.isna().sum()),'filas_idps':len(idps),
    'oc_seleccionadas':len(oc),'oc_con_rbd':len(linked),'colegios_prioridad_50':len(priority),
    'colegios_sin_pme':len(without_pme),'dimension_mayor_kpi':highest.dimension,
    'kpi_dimension_mayor':float(highest.pct_senal_revision),'validaciones':len(validation)}
(OUT/'resultados_eda.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
score24=simce.loc[simce.anio.eq(2024)&simce.grado.eq('4b')&simce.area.eq('lect'),'puntaje']
cc24=idps.loc[idps.anio.eq(2024)&idps.grado.eq('4b')&idps.indicador_std.eq('CC'),'puntaje']
note(f'**Contexto educativo concreto:** Lectura 4b 2024 tiene {score24.notna().sum()} puntajes, '
     f'mediana {score24.median():.1f} e IQR {score24.quantile(.75)-score24.quantile(.25):.1f}. '
     f'Clima de convivencia 4b 2024 tiene {cc24.notna().sum()} puntajes, mediana {cc24.median():.1f} '
     f'e IQR {cc24.quantile(.75)-cc24.quantile(.25):.1f}. Son distribuciones entre establecimientos, no alumnos.')
note(f'''**Respuesta:** en las {len(pme)} acciones observadas de {len(schools)} colegios, {pme.senal_revision.sum()} ({100*pme.senal_revision.mean():.1f}%) tienen estimación cero o no se declaran completas. La mayor proporción por dimensión aparece en **{highest.dimension}**, con {highest.pct_senal_revision:.1f}%. Con el umbral propuesto de 50%, {len(priority)} colegios entrarían primero a revisión documental.

**Importes:** la mediana estimada por acción es ${esnum(pme.estimacion_clp.median())}, frente a una media de ${esnum(pme.estimacion_clp.mean())}. El 10% superior concentra {100*topshare:.1f}% de la estimación. La cola y los {pme.extremo_iqr.sum()} casos marcados por IQR justifican revisar escala y documentación, sin borrar observaciones.

**Disponibilidad:** faltan {simce.puntaje.isna().sum()} de {len(simce)} puntajes SIMCE por área dentro de filas publicadas y {idps.puntaje.isna().sum()} de {len(idps)} promedios IDPS. El panel además tiene ausencias de fila y bloques fuera del período recopilado. Las matrices muestran ambas situaciones por separado.

**Compras:** {len(oc)} órdenes seleccionadas y {len(linked)} vinculadas a un único RBD. La deduplicación evita repetir el total por ítem. La cobertura desigual de 2022 y 2025 impide leer su diferencia como una tendencia de gasto escolar.

**Asociaciones:** las matrices Pearson/Spearman, sus tamaños por pareja y la segmentación permiten explorar dependencia entre variables. Ningún coeficiente establece que el monto PME, una OC o el avance declarado causen un resultado SIMCE/IDPS.

**Alcance de la decisión:** priorizar documentación es viable; cuantificar malgasto o eficacia causal requiere otros datos. Los {len(without_pme)} colegios sin PME no reciben un cero ni una etiqueta de incumplimiento.''')

# %% [markdown]
# ### 12.1 Hallazgos defendibles frente a provisionales
#
# La pauta del curso pide distinguir lo que el EDA **permite afirmar** de lo que solo **sugiere como hipótesis**. La tabla se construye con las cifras calculadas arriba, de modo que se actualiza si se repite la ejecución.

# %%
oc_factor=(oc_audit.suma_incorrecta_por_fila_clp/oc_audit.suma_correcta_por_oc_clp)[oc_audit.anio.le(2024)]
ok_checks=int(pd.DataFrame(validation).resultado.eq('OK').sum())
findings=pd.DataFrame([
 ('Defendible','Los datos usados son trazables y consistentes',f'{ok_checks} de {len(validation)} validaciones OK; 0 filas eliminadas en la limpieza; huellas SHA-256 verificadas','Verifica el extracto, no la veracidad de lo declarado por los colegios'),
 ('Defendible','Sumar la OC por ítem sobreestima el monto',f'Factor de {oc_factor.min():.0f} a {oc_factor.max():.0f} veces en 2022–2024 frente a contar cada código una vez','Una OC no acredita pago'),
 ('Defendible','Los importes PME son muy asimétricos',f'Media ${esnum(pme.estimacion_clp.mean())} vs. mediana ${esnum(pme.estimacion_clp.median())}; el 10 % superior concentra {100*topshare:.1f} % del total','Se conservan los extremos; su efecto se mide en la sensibilidad'),
 ('Defendible','Casi la mitad de las acciones PME 2024 requieren explicación',f'{int(pme.senal_revision.sum())} de {len(pme)} acciones ({100*pme.senal_revision.mean():.1f} %) con $0 o no completas','Depende de la definición: con la estricta baja a '+f'{100*pme.senal_estricta.mean():.1f} %'),
 ('Defendible','El KPI por colegio es bimodal',f'{low} colegios ≤25 %, {mid} intermedios, {high} ≥75 %','Describe el extracto de 46 colegios'),
 ('Defendible','Las ausencias son estructurales; imputar las distorsionaría','PME solo 2024, Agencia desde 2023, OC solo del sector público (ver mapa de faltantes)','Un cero válido y un dato ausente se reportan por separado'),
 ('Provisional','El umbral de 50 % sirve para priorizar',f'Entran {len(priority)} de {len(schools)} colegios; con 25 % y 75 % cambia a {int(schools.pct_senal_revision.ge(25).sum())} y {int(schools.pct_senal_revision.ge(75).sum())}','Supuesto operativo sin validar con capacidad real de revisión'),
 ('Provisional','Las escuelas especiales concentran las estimaciones $0',f'{zero_by_offer["especial"]:.1f} % vs. {zero_by_offer["basica_media"]:.1f} % en básica/media','Hipótesis sobre su financiamiento; requiere confirmar con Mineduc o el sostenedor'),
 ('Provisional','Colegios pequeños tienden a KPI más alto',f'Spearman KPI–matrícula = {rho_size:.2f}','Posible efecto de escala; no causal; n pequeño y segmento público con n<10'),
 ('Provisional','No se observa relación entre KPI y resultados SIMCE/IDPS','Correlaciones cercanas a 0 con n entre 28 y 29 colegios','Poca potencia; los resultados dependen del grupo evaluado'),
 ('Provisional','Las compras 2022 y 2025 no son comparables con 2023–2024','Falta Compra Ágil del 2.º semestre 2022 y la cobertura SLEP 2025 es parcial','No se interpreta una caída como ahorro'),
 ('Provisional','Los 13 colegios sin PME están fuera del KPI',f'{len(without_pme)} RBD sin acciones en el extracto','Verificar aplicabilidad y cobertura antes de cualquier conclusión'),
],columns=['estatus','hallazgo','evidencia en este notebook','límite o siguiente paso'])
table(findings,'hallazgos_defendibles_vs_provisionales')

# %% [markdown]
# ### 12.2 Los siete pasos de la clase 03 aplicados al proyecto
#
# Los nombres siguen el brief de C1. El usuario y la decisión son una **propuesta**: no hubo entrevistas ni validación con actores externos.
#
# | Paso | Qué se responde en este proyecto | Evidencia |
# |---|---|---|
# | 1. Decisión | Usuarios propuestos: equipos directivos, sostenedores y equipo investigador. Decisión: qué respaldo del PME solicitar primero. Alternativas: pedir respaldo de todas las acciones, ordenar por monto, por resultados SIMCE o por sector. Se eligió el KPI porque conecta una señal observable con un documento concreto | Secciones 5.1 y 11.4 |
# | 2. Criterios de éxito | Deseado: que cada solicitud apunte a acciones cuyo registro necesita explicación. Medible hoy: proporción de acciones con Z o N y número de colegios sobre un umbral. No medible hoy: si la señal correspondía a un problema real, porque no hay etiquetas verificadas | Secciones 5.1, 11.3 y 12.1 |
# | 3. Preguntas | Principal: cómo se distribuye el KPI entre colegios y dimensiones y qué casos priorizar. De apoyo: cuánto cambia la lista al variar la definición y el umbral; si la señal se asocia con tamaño, sector o resultados; qué falta y por qué | Portada, secciones 8, 10 y 11.3 |
# | 4. Proceso | Planificación anual, Implementación (registro del estado y ajuste de montos), Evaluación, rendición y fiscalización. Los datos capturan Planificación e Implementación. **No capturan** rendiciones, pagos, recepción, medios de verificación ni la Evaluación | Sección 1 y 11.4 |
# | 5. Entidades, eventos y estados | Establecimiento (RBD), acción PME, dimensión, estado declarado de implementación, estimación en CLP, orden de compra, proceso PAS y medida SIMCE o IDPS por grado. Un colegio no es una acción y una acción no es una orden de compra | Secciones 1 y 5.2 |
# | 6. Representación de datos | Una fila es: una acción (PME), un RBD–año (Directorio y panel), un RBD–año–grado–área (SIMCE), un RBD–año–grado–indicador (IDPS), una orden de compra o un proceso PAS. Los cruces son por RBD y año y siempre después de agregar | Secciones 5.2 y 6.1 |
# | 7. Medición | KPI = 100 × acciones con Z o N ÷ acciones observadas en 2024; unidad: porcentaje; numerador y denominador siempre reportados; período 2024; exclusión: 13 colegios sin acciones; límite: no mide malgasto ni calidad | Secciones 5.1 y 11 |

# %% [markdown]
# ### 12.3 Las Five Cs aplicadas a esta investigación
#
# Nombres del brief de C1: Clean, Consistent, Conformed, Current y Comprehensive. Las cifras se calculan en la celda siguiente. Ningún hallazgo se inventa para parecer más sofisticado: donde no hubo problema, se reporta la comprobación.

# %%
neg=int(pme.estimacion_clp.lt(0).sum()); sep_gt=int(pme.sep_clp.gt(pme.estimacion_clp).sum())
idps_out=int((idps.puntaje.lt(0)|idps.puntaje.gt(100)).sum())
dup_pme=int(pme.duplicated(['RBD','DIMENSION','ESTIM_TOTAL','ESTIM_SEP','NIV_IMPLEM','NOM_ACTIVIDAD']).sum())
codes=sorted(pme.COD_DEPE.dropna().astype(int).unique().tolist())
mismatch=int(((pme.COD_DEPE.isin([2,6])&pme.sector.ne('Público'))|(pme.COD_DEPE.eq(3)&pme.sector.ne('Particular subvencionado'))).sum())
grades={int(k):v for k,v in simce.dropna(subset=['puntaje']).groupby('anio').grado.agg(lambda g:', '.join(sorted(set(g)))).items()}
grades_txt='; '.join(f'{k}: {v}' for k,v in grades.items())
public_dep=directory.loc[directory.sector.eq('Público')].groupby('anio').dependencia.agg(lambda g:', '.join(sorted(set(g)))).to_dict()
idps_numeric=int(idps.indicador_codigo_original.astype(str).isin(['1','2','3','4']).sum())
s4=int(simce.loc[simce.anio.eq(2024)&simce.grado.eq('4b')&simce.area.eq('lect')&simce.puntaje.notna(),'rbd'].nunique())
five_c=pd.DataFrame([
 ('Clean','Validez y duplicados',
  f'{dup_pme} acciones PME duplicadas; {neg} importes negativos; {sep_gt} con SEP mayor que el total; {idps_out} promedios IDPS fuera de 0–100; llaves únicas en SIMCE, IDPS, OC y PAS; {len(validation)} validaciones superadas',
  'No se corrigió ningún valor ni se eliminó ninguna fila. Se conservan ceros y extremos (se marcan)','Secciones 3, 4 y 6'),
 ('Consistent','Estabilidad de las definiciones dentro de cada fuente entre períodos',
  f'IDPS cambia sus códigos en 2025 ({idps_numeric} filas con código numérico, homologadas por nombre); SIMCE evalúa grados distintos por año ({grades_txt}; 2m = 2.º medio, 4b = 4.º básico, 6b = 6.º básico, 8b = 8.º básico); los colegios públicos pasan de «{public_dep[2024]}» en 2024 a «{public_dep[2025]}» en 2025; el PME usa un único vocabulario de estados en 2024',
  'Se homologa IDPS, no se mezclan grados ni aplicaciones y el PME se analiza en un solo corte','Secciones 4, 7.3 y 7.4'),
 ('Conformed','Acuerdo de entidades y códigos entre las fuentes que se combinan',
  f'RBD común: todas las filas SIMCE e IDPS pertenecen al censo del año; dependencia en PME (códigos {codes}) frente al sector del Directorio: {mismatch} discordancias en {len(pme)} acciones; nombres de indicador IDPS únicos por código homologado',
  'Los cruces son por RBD y año; una OC se asigna solo si menciona un único RBD','Secciones 6 y 6.1'),
 ('Current','Antigüedad y disponibilidad frente a la decisión',
  'La Implementación 2024 cerró el 14-12-2024 y la Evaluación el 15-01-2025 (aviso de Mineduc del 03-01-2025); a septiembre de 2026 los datos tienen casi dos años y el registro es dinámico; las compras 2025 son parciales',
  'Sirven para ordenar una revisión documental retrospectiva; no para monitoreo actual. Repetir el análisis con otros años es un paso siguiente','Secciones 1 y 7.2'),
 ('Comprehensive','Cobertura de población, campos, períodos y excepciones',
  f'{len(schools)} de {len(census24)} colegios del censo 2024 tienen acciones PME y {len(without_pme)} no; SIMCE 4.º básico Lectura 2024: {s4} colegios; compras solo de los 8 colegios públicos; PAS sin materia de cargos. No observado: rendiciones, pagos, recepción y medios de verificación',
  'Lo no observado no se rellena; los colegios sin acciones no reciben KPI cero','Secciones 8, 8.2 y 12')],
 columns=['C','Qué se comprobó','Evidencia calculada','Decisión y límite','Dónde'])
table(five_c,'cinco_c_aplicadas')

# %% [markdown]
# ### 12.4 Continuidad: tarea de ML y dashboard posibles (propuesta, no implementada)
#
# Esta sección es una **propuesta para revisión**. No se entrenó ningún modelo ni se construyó el dashboard, y la viabilidad no está demostrada.
#
# **Restricción de partida.** No hay etiquetas verificadas: no se sabe qué acciones o colegios tenían un problema real. El KPI es una señal para pedir explicaciones, no una verdad comprobada, y no debe usarse como si lo fuera.
#
# **Tarea A, exploratoria y la más viable con los datos actuales: perfiles de colegios con clustering.**
#
# - Unidad de análisis: colegio (RBD) en 2024, 46 casos.
# - Entradas disponibles: proporción de acciones por dimensión, mediana de la estimación, número de acciones, matrícula, sector y oferta.
# - No se define ninguna etiqueta.
# - Valor para la decisión: identificar tipos de colegio para ordenar la cola de revisión y pedir respaldos distintos según el perfil.
# - Incertidumbres para la etapa siguiente: 46 casos es una muestra pequeña, las variables tienen escalas muy distintas y el resultado puede depender del número de grupos.
#
# **Tarea B, supervisada y condicionada a obtener más datos: anticipar si una acción terminará sin implementación completa.**
#
# - Unidad de análisis: acción PME.
# - Resultado propuesto: estado declarado al cierre de la Implementación (variable N).
# - Información disponible al planificar: dimensión, estimación total y SEP, sector, oferta y matrícula.
# - Información que solo existe después del resultado y no puede usarse como entrada: estado declarado, avance y ajustes posteriores del monto.
# - Advertencias: el $0 (variable Z) forma parte del KPI y no debe ser a la vez entrada y resultado; N incluye avances de 75 a 99 %, por lo que el resultado es ruidoso.
# - Datos adicionales necesarios: PME de otros años (el portal publica Planificación e Implementación desde 2023) para separar entrenamiento y prueba en el tiempo.
# - Valor para la decisión: anticipar dónde acompañar antes de que cierre el año.
# - Incertidumbres: hoy hay un solo año extraído, las estimaciones se ajustan durante la implementación y podría convenir otra definición del resultado.
#
# **Dashboard.**
#
# - Usuarios propuestos: quien solicita respaldos, ya sea un sostenedor o el equipo investigador.
# - Decisiones: qué colegio y qué acciones revisar primero.
# - Indicadores centrales: KPI con numerador y denominador, acciones con $0, acciones no completas, mediana de la estimación y colegios sobre el umbral.
# - Comparaciones y filtros: dimensión, sector, oferta, umbral (25, 50 y 75 %) y colegio, con detalle hacia la lista de acciones (nombre, estimación y estado).
# - Cómo lo informa el EDA: la distribución bimodal justifica ordenar por colegio; la sensibilidad a la definición exige mostrar la variante estricta junto a la principal; los faltantes exigen mostrar aparte los 13 colegios sin acciones.

# %% [markdown]
# ## 13. Reproducibilidad y trazabilidad de la entrega
#
# Este cuaderno conserva datos originales, reglas y salidas en secuencia. Las figuras tienen denominadores, unidades, poblaciones y escalas identificadas. Se guardan inventario individual, perfiles, balances, validaciones, tablas de faltantes, extremos, matrices de correlación con tamaños y listados del KPI. Todas las tablas se derivan en este mismo notebook.
#
# El pipeline base construye `panel_rbd_anual.csv`. Este notebook crea su extensión analítica `eda/panel_rbd_anual_eda.csv`. La lectura comprimida de OC está en `analysis/scripts/integrar_chilecompra.py`; sus decisiones y efectos quedan documentados arriba. Los indicadores y gráficos no dependen de cálculos ocultos en un informe.
#
# La copia HTML `analysis/EDA_gasto_educativo.html` facilita leer las mismas salidas sin Jupyter; no es otro análisis. El registro humano/IA, sugerencias adaptadas, verificaciones y pendientes se conserva en [PROCESS_LOG.md](../PROCESS_LOG.md). No se inventan contribuciones personales ni resultados de entrevistas.

# %% [markdown]
# ### 13.1 Datos para la presentación, derivados de las mismas tablas
#
# Se exportan series editables para las diapositivas. Las clases de histogramas y las densidades se calculan desde las mismas observaciones, con sus filtros declarados. Esto evita transcribir cifras o usar una captura como sustituto de datos editables.

# %%
def histogram_payload(s,bins):
    values=pd.Series(s).dropna().to_numpy(dtype=float)
    counts,edges=np.histogram(values,bins=bins)
    return {'n':len(values),'centros':((edges[:-1]+edges[1:])/2).tolist(),
            'limites':edges.tolist(),'frecuencias':counts.tolist()}
def kde_payload(s,points=40):
    values=pd.Series(s).dropna().to_numpy(dtype=float)
    grid=np.linspace(values.min(),values.max(),points)
    return {'n':len(values),'x':grid.tolist(),'y':gaussian_kde(values,bw_method='scott')(grid).tolist()}
payload={'resultados':result,'dimensiones':dimensions.to_dict('records'),
         'oc_antes_despues':oc_audit.to_dict('records'),'umbrales':thresholds.to_dict('records'),
         'sensibilidad':sensitivity.to_dict('records'),'top_colegios':priority.head(8).to_dict('records'),
         'pme_hist':histogram_payload(pme.estimacion_clp/1e6,np.linspace(0,250,11)),
         'pme_kde':kde_payload(np.log10(pme.loc[pme.estimacion_clp.gt(0),'estimacion_clp'])),
         'oc_hist_2024':histogram_payload(np.log10(oc_included.loc[oc_included.anio.eq(2024),'monto_oc_vigente_clp']),np.linspace(4,8,11)),
         'simce_hist_4b2024':histogram_payload(score24,np.arange(180,341,20)),
         'idps_hist_cc4b2024':histogram_payload(cc24,np.arange(40,101,5)),
         'simce_kde_4b2024':kde_payload(score24),'idps_kde_cc4b2024':kde_payload(cc24),
         'corr_vars':corr_vars,'corr_names':corr_names,'spearman':spearman.where(spearman.notna(),None).to_numpy().tolist(),
         'pearson':pearson.where(pearson.notna(),None).to_numpy().tolist(),'corr_n':counts.to_numpy().tolist(),
         'missing_panel':panel.groupby('anio')[missing_cols].agg(lambda s:float(100*s.isna().mean())).reset_index().to_dict('records'),
         'missing_agencia':simce_missing.to_dict('records'),'missing_idps':idps_missing.to_dict('records'),
         'asociaciones':pd.DataFrame(assoc).to_dict('records')}
# JSON estricto: las correlaciones no calculables se serializan como null.
payload=json.loads(pd.Series({'payload':payload}).to_json(force_ascii=False))['payload']
(OUT/'presentacion_datos.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
print('Series de presentación exportadas desde el notebook.')

# %% [markdown]
# ## 14. Anexo para defensa: siete pasos y Five Cs
#
# Los siete pasos (sección 12.2) y las Five Cs (sección 12.3) se aplican arriba; la tabla siguiente conserva el mapeo original de W1.
#
# | Siete pasos de W1 | Aplicación revisada |
# |---|---|
# | Decisión | Elegir qué respaldo PME solicitar primero |
# | Éxito | Lista trazable con numerador, denominador y motivo de revisión |
# | Pregunta de negocio | Distribución de acciones con cero o no completas en 2024 |
# | Proceso | Registro PME, revisión, solicitud, cotejo y cierre documental |
# | Entidades/eventos/estados | Colegio, acción, dimensión, estado, estimación, documento |
# | Representación | Acción para señales; RBD 2024 para priorización; panel contextual |
# | Medición | KPI unión Z o N y sus componentes, cobertura y sensibilidad |
#
# **Five Cs (nombres oficiales del brief de C1):** Clean, Consistent, Conformed, Current y Comprehensive. Se aplican con evidencia calculada en la sección 12.3. Los nombres usados en W1 (Complete, Correct, Consistent, Currency y Concordant/Unique) eran provisionales y no coinciden con los del brief.
#
# **Entrega de C1:** según el brief, un archivo ZIP `GROUP_03_C1.zip` por Canvas con `README.md`, `report/`, `analysis/`, `data/` y `presentation/`; GitHub es opcional. La calendarización del curso menciona repositorio y SHA; el brief de C1, más específico, prevalece.
#
# **Materiales contrastados:** calendarización recuperada; `03_EDA_after_Data_Validation.pdf` (distribuciones, IQR/MAD, correlaciones, segmentación e interpretación); notebooks de clase disponibles; W1 con reflexiones; presentación previa. No se encontró un archivo independiente titulado «pseudo temario»; sus exigencias mencionadas por el usuario se aplican mediante el flujo y contenidos enumerados, sin inventar una pauta adicional.
#
# **Fuentes oficiales:** [Directorio Mineduc](https://datosabiertos.mineduc.cl/directorio-de-establecimientos-educacionales/), [PME](https://liderazgoeducativo.mineduc.cl/bases-de-datos-pme/), [PAS Supereduc](https://www.supereduc.cl/pas/), [Agencia de Calidad](https://informacionestadistica.agenciaeducacion.cl/), [ChileCompra](https://datos-abiertos.chilecompra.cl/descargas/ordenes-y-licitaciones). Las descargas, versiones y fechas concretas están en los manifiestos locales. No se reemplazaron archivos por versiones web nuevas durante esta revisión.

# %%
artifacts=sorted(p.name for p in OUT.glob('*.csv'))
note(f'**Cierre de ejecución:** {len(validation)} validaciones sustantivas superadas; '
     f'{len(artifacts)} tablas CSV disponibles en la carpeta `eda/`. '
     'Los gráficos visibles son salidas de esta ejecución.')
table(pd.DataFrame({'tabla_exportada':artifacts}))
